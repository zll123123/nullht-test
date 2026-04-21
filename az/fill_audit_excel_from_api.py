#!/usr/bin/env python3
import argparse
import json
import ssl
import sys
import urllib.request
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from openpyxl import load_workbook


DEFAULT_BASE_URL = "https://dev-api-v3-az-mlr.nullht.com"
DEFAULT_LIST_PATH = "/api/audit/management/list"
DEFAULT_DETAIL_PATH = "/api/audit/management/detail"
DEFAULT_COOKIE = "acw_tc=65859a8117763273964628718ece2b4a8be2d51449edc804cb8ba1237c0fc9"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="根据 Excel 中的任务编号调用审核管理接口，并将接口结果回填到 Excel。")
    parser.add_argument("--excel-path", required=True, help="待回填的 Excel 文件路径")
    parser.add_argument("--bearer-token", required=True, help="管理接口 Bearer token，传入纯 token 即可")
    parser.add_argument("--cookie", default=DEFAULT_COOKIE, help="请求使用的 Cookie")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help="审核管理接口域名")
    parser.add_argument("--sheet-name", help="可选，指定工作表名称；默认使用 active sheet")
    parser.add_argument(
        "--keep-existing-rows",
        action="store_true",
        help="默认会删除第 2 行之后的旧数据并重写；传入该参数后将只追加，不删除旧数据",
    )
    return parser.parse_args()


def build_headers(token: str, cookie: str) -> Dict[str, str]:
    return {
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh;q=0.9",
        "Authorization": f"Bearer {token}",
        "Connection": "keep-alive",
        "Content-Type": "application/json",
        "Cookie": cookie,
        "Origin": "https://dev-v3-az-mlr.nullht.com",
        "Referer": "https://dev-v3-az-mlr.nullht.com/",
        "Sec-Fetch-Dest": "empty",
        "Sec-Fetch-Mode": "cors",
        "Sec-Fetch-Site": "same-site",
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/147.0.0.0 Safari/537.36",
        "sec-ch-ua": '"Google Chrome";v="147", "Not.A/Brand";v="8", "Chromium";v="147"',
        "sec-ch-ua-mobile": "?0",
        "sec-ch-ua-platform": '"macOS"',
    }


def post_json(url: str, payload: Dict[str, Any], headers: Dict[str, str]) -> Dict[str, Any]:
    request = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    context = ssl.create_default_context()
    with urllib.request.urlopen(request, context=context, timeout=30) as response:
        return json.loads(response.read().decode("utf-8"))


def load_seed_rows(sheet) -> List[Tuple[str, str]]:
    rows: List[Tuple[str, str]] = []
    seen: Set[str] = set()
    for row_index in range(2, sheet.max_row + 1):
        task_id = sheet.cell(row_index, 1).value
        file_name = sheet.cell(row_index, 2).value
        if not task_id:
            continue
        task_id = str(task_id).strip()
        if task_id in seen:
            continue
        rows.append((task_id, "" if file_name is None else str(file_name)))
        seen.add(task_id)
    return rows


def clear_data_rows(sheet) -> None:
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)


def fetch_list_row(base_url: str, task_id: str, headers: Dict[str, str]) -> Optional[Dict[str, Any]]:
    response = post_json(
        f"{base_url}{DEFAULT_LIST_PATH}",
        {
            "page_num": 1,
            "page_size": 10,
            "status": [],
            "task_id": task_id,
            "sorts": ["CREATED_TIME_DESC"],
        },
        headers,
    )
    if response.get("code") != "00000":
        raise RuntimeError(f"LIST 返回异常: {response.get('code')} {response.get('msg')}")
    return ((response.get("data") or {}).get("rows") or [None])[0]


def fetch_detail(base_url: str, row_id: str, headers: Dict[str, str]) -> Dict[str, Any]:
    response = post_json(f"{base_url}{DEFAULT_DETAIL_PATH}", {"id": row_id}, headers)
    if response.get("code") != "00000":
        raise RuntimeError(f"DETAIL 返回异常: {response.get('code')} {response.get('msg')}")
    return response.get("data") or {}


def expand_findings(task_id: str, file_name: str, detail: Dict[str, Any]) -> List[Dict[str, Any]]:
    page_details = detail.get("page_details") or []
    if not page_details:
        return [
            {
                "task_id": task_id,
                "file_name": file_name,
                "page_number": "无",
                "point_name": "无",
                "reason": "detail.page_details 为空",
            }
        ]

    rows: List[Dict[str, Any]] = []
    for page_detail in page_details:
        default_page_number = page_detail.get("page_number")
        for finding in page_detail.get("findings") or []:
            positions = finding.get("position") or []
            page_number = default_page_number
            if positions and positions[0].get("page_number") is not None:
                page_number = positions[0]["page_number"]
            rows.append(
                {
                    "task_id": task_id,
                    "file_name": file_name,
                    "page_number": page_number,
                    "point_name": finding.get("point_name") or "",
                    "reason": finding.get("reason")
                    or ((finding.get("extension") or {}).get("llm_reason"))
                    or "",
                }
            )
    return rows


def write_rows(sheet, rows: List[Dict[str, Any]], keep_existing_rows: bool) -> None:
    start_row = sheet.max_row + 1 if keep_existing_rows and sheet.max_row >= 2 else 2
    for index, row in enumerate(rows, start=start_row):
        sheet.cell(index, 1).value = row["task_id"]
        sheet.cell(index, 2).value = row["file_name"]
        sheet.cell(index, 3).value = row["page_number"]
        sheet.cell(index, 4).value = row["point_name"]
        sheet.cell(index, 5).value = row["reason"]


def main() -> int:
    args = parse_args()
    excel_path = Path(args.excel_path).expanduser().resolve()
    workbook = load_workbook(excel_path)
    sheet = workbook[args.sheet_name] if args.sheet_name else workbook.active
    seed_rows = load_seed_rows(sheet)
    headers = build_headers(args.bearer_token, args.cookie)

    if not args.keep_existing_rows:
        clear_data_rows(sheet)

    output_rows: List[Dict[str, Any]] = []
    for task_id, file_name in seed_rows:
        try:
            list_row = fetch_list_row(args.base_url, task_id, headers)
            if not list_row:
                output_rows.append(
                    {
                        "task_id": task_id,
                        "file_name": file_name,
                        "page_number": "未查询到审核记录",
                        "point_name": "未查询到审核记录",
                        "reason": "未查询到审核记录",
                    }
                )
                continue
            detail = fetch_detail(args.base_url, list_row["id"], headers)
            output_rows.extend(expand_findings(task_id, file_name, detail))
        except Exception as exc:
            output_rows.append(
                {
                    "task_id": task_id,
                    "file_name": file_name,
                    "page_number": f"处理失败: {exc}",
                    "point_name": f"处理失败: {exc}",
                    "reason": f"处理失败: {exc}",
                }
            )

    output_rows.sort(key=lambda item: (item["task_id"], str(item["page_number"]), item["point_name"]))
    write_rows(sheet, output_rows, args.keep_existing_rows)
    workbook.save(excel_path)

    print(f"Excel 已更新: {excel_path}")
    print(f"写入记录数: {len(output_rows)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
