from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import requests


EXCEL_PATH = Path("/Users/layla.zhang/workspace/nullht-test/az/拉取线上原始文件/结果统计base_docx.xlsx")
SHEET_NAME = "明细"
LIST_URL = "https://dev-api-v3-az-mlr.nullht.com/api/audit/management/list"
TIMEOUT = 30
AUTHORIZATION = (
    "Bearer eyJ0eXAiOiJKV1QiLCJhbGciOiJIUzI1NiJ9.eyJsb2dpblR5cGUiOiJsb2dpbiIsImxvZ2luSWQiOiIwOGY4ZjU5NTQ4ZDQ0ODg4Yjk4ZTQ5OTkwYjUxMDQ5ZSIsImRldmljZVR5cGUiOiJERUYiLCJlZmYiOjE3Nzk0NDU2NTExMDIsInJuU3RyIjoic2I1Z3p5RHhoZDBYdDFsZlhnbHZkTnNhS1BybVVnejMiLCJhel9vcGVuX2lkIjoieHh4eCJ9.gf3zDhkHNSDgrDV_SG30uRjtvAv-3zm5oHNShb5RLp0"
)
COOKIE = "acw_tc=65859a8117793605590861657eca1f38b24c158300d3e9e7c"

DETAIL_NEW_COLUMN = "detail_new"
TASK_ID_COLUMN = "task_id"
FILE_NAME_COLUMN = "file_name"


def build_headers() -> dict[str, str]:
    """构造请求头。"""

    return {
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh;q=0.9",
        "Authorization": AUTHORIZATION,
        "Connection": "keep-alive",
        "Content-Type": "application/json",
        "Cookie": COOKIE,
        "Origin": "https://dev-v3-az-mlr.nullht.com",
        "Referer": "https://dev-v3-az-mlr.nullht.com/",
        "Sec-Fetch-Dest": "empty",
        "Sec-Fetch-Mode": "cors",
        "Sec-Fetch-Site": "same-site",
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36"
        ),
        "sec-ch-ua": '"Chromium";v="148", "Google Chrome";v="148", "Not/A)Brand";v="99"',
        "sec-ch-ua-mobile": "?0",
        "sec-ch-ua-platform": '"macOS"',
    }


def find_column(worksheet: Worksheet, column_name: str) -> int:
    """查找表头列号。"""

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    raise ValueError(f"Excel 缺少表头: {column_name}")


def get_or_create_column(worksheet: Worksheet, column_name: str) -> int:
    """获取或创建表头列。"""

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    new_column_index = worksheet.max_column + 1
    worksheet.cell(row=1, column=new_column_index, value=column_name)
    return new_column_index


def get_cell_value(worksheet: Worksheet, row_index: int, column_index: int) -> str:
    """读取单元格文本。"""

    cell_value = worksheet.cell(row=row_index, column=column_index).value
    return "" if cell_value is None else str(cell_value).strip()


def fetch_first_result(session: requests.Session, file_name: str) -> tuple[str, str]:
    """根据 file_name 获取 list 首条 id 和 task_id。"""

    payload = {
        "page_num": 1,
        "page_size": 10,
        "status": [],
        "file_name": file_name,
        "sorts": ["CREATED_TIME_DESC"],
    }
    response = session.post(LIST_URL, json=payload, timeout=TIMEOUT)
    response.raise_for_status()
    response_json = response.json()
    if str(response_json.get("code", "")).strip() != "00000":
        raise RuntimeError(
            f"list接口返回异常: code={response_json.get('code')}, msg={response_json.get('msg')}"
        )
    rows = ((response_json.get("data") or {}).get("rows") or [])
    if not rows:
        return "", ""
    first_row = rows[0] or {}
    detail_id = str(first_row.get("id", "")).strip()
    task_id = str(first_row.get("task_id", "")).strip()
    return detail_id, task_id


def fill_grouped_detail_and_task_id(
    worksheet: Worksheet,
    file_name_column: int,
    detail_new_column: int,
    task_id_column: int,
) -> None:
    """按连续相同 file_name 分组回填 detail_new 和 task_id。

    Args:
        worksheet: Excel 工作表。
        file_name_column: file_name 列号。
        detail_new_column: detail_new 列号。
        task_id_column: task_id 列号。
    """

    last_file_name = ""
    last_detail_id = ""
    last_task_id = ""

    for row_index in range(2, worksheet.max_row + 1):
        file_name = get_cell_value(worksheet, row_index, file_name_column)
        detail_id = get_cell_value(worksheet, row_index, detail_new_column)
        task_id = get_cell_value(worksheet, row_index, task_id_column)

        if file_name != last_file_name:
            last_file_name = file_name
            last_detail_id = detail_id
            last_task_id = task_id
            continue

        if not detail_id and last_detail_id:
            worksheet.cell(row=row_index, column=detail_new_column, value=last_detail_id)
        if not task_id and last_task_id:
            worksheet.cell(row=row_index, column=task_id_column, value=last_task_id)

        if detail_id:
            last_detail_id = detail_id
        if task_id:
            last_task_id = task_id


def main() -> None:
    """执行回填。"""

    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook[SHEET_NAME]
    detail_new_column = find_column(worksheet, DETAIL_NEW_COLUMN)
    task_id_column = get_or_create_column(worksheet, TASK_ID_COLUMN)
    file_name_column = find_column(worksheet, FILE_NAME_COLUMN)
    session = requests.Session()
    session.headers.update(build_headers())
    file_name_cache: dict[str, tuple[str, str]] = {}

    for row_index in range(2, worksheet.max_row + 1):
        file_name = get_cell_value(worksheet, row_index, file_name_column)
        if not file_name:
            continue
        if file_name in file_name_cache:
            detail_id, task_id = file_name_cache[file_name]
        else:
            detail_id, task_id = fetch_first_result(session=session, file_name=file_name)
            file_name_cache[file_name] = (detail_id, task_id)
        worksheet.cell(row=row_index, column=detail_new_column, value=detail_id)
        worksheet.cell(row=row_index, column=task_id_column, value=task_id)
        print(
            f"回填完成: row={row_index}, file_name={file_name}, "
            f"detail_new={detail_id}, task_id={task_id}"
        )

    fill_grouped_detail_and_task_id(
        worksheet=worksheet,
        file_name_column=file_name_column,
        detail_new_column=detail_new_column,
        task_id_column=task_id_column,
    )
    workbook.save(EXCEL_PATH)


if __name__ == "__main__":
    main()
