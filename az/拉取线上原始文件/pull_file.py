from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import requests


EXCEL_PATH = Path("/Users/layla.zhang/workspace/nullht-test/az/拉取线上原始文件/结果统计base.xlsx")
SHEET_NAME = "明细"
DOWNLOAD_DIR = Path("/Users/layla.zhang/测试用例/测试材料/az/产品图片原始文件")
DETAIL_URL = "https://dev-api-v3-az-mlr.nullht.com/api/audit/management/detail"
DOWNLOAD_URL = "https://dev-api-v3-az-mlr.nullht.com/api/oss/download"
TIMEOUT = 30
AUTHORIZATION = (
    "Bearer "
    "eyJ0eXAiOiJKV1QiLCJhbGciOiJIUzI1NiJ9."
    "eyJsb2dpblR5cGUiOiJsb2dpbiIsImxvZ2luSWQiOiIwOGY4ZjU5NTQ4ZDQ0ODg4Yjk4ZTQ5OTkwYjUxMDQ5ZSIs"
    "ImRldmljZVR5cGUiOiJERUYiLCJlZmYiOjE3NzkxNjMzMTQwNzgsInJuU3RyIjoiNlpBdUVVVk9xZWI0R1NCUkh4"
    "VjJNR29QWTVjU0J2N1kiLCJhel9vcGVuX2lkIjoieHh4eCJ9."
    "IYy3YeZFD6Nm074NB73VmMy5eHjUE0xsxVqAJPGGZxw"
)
COOKIE = "acw_tc=65859a8117790769001861223ece098c6748d9336d91bb20d51ee0b21a09b0"

FILE_ID_COLUMN = "download_file_id"
FILE_NAME_COLUMN = "file_name"
DOWNLOAD_STATUS_COLUMN = "download_status"
DOWNLOAD_PATH_COLUMN = "download_path"
DOWNLOAD_ERROR_COLUMN = "download_error"


def get_or_create_column(worksheet: Worksheet, column_name: str) -> int:
    """获取或创建表头列。

    Returns:
        int: 列号。
    """

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    new_column_index = worksheet.max_column + 1
    worksheet.cell(row=1, column=new_column_index, value=column_name)
    return new_column_index


def build_headers() -> dict[str, str]:
    """构造请求头。

    Returns:
        dict[str, str]: 请求头字典。
    """

    headers = {
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh;q=0.9",
        "Connection": "keep-alive",
        "Content-Type": "application/json",
        "Origin": "https://dev-v3-az-mlr.nullht.com",
        "Referer": "https://dev-v3-az-mlr.nullht.com/",
        "Sec-Fetch-Dest": "empty",
        "Sec-Fetch-Mode": "cors",
        "Sec-Fetch-Site": "same-site",
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36"
        ),
        "sec-ch-ua": '"Chromium";v="148", "Google Chrome";v="148", "Not/A)Brand";v="99"',
        "sec-ch-ua-mobile": "?0",
        "sec-ch-ua-platform": '"macOS"',
    }
    authorization = normalize_authorization(AUTHORIZATION)
    if authorization:
        headers["Authorization"] = authorization
    cookie = normalize_cookie(COOKIE)
    if cookie:
        headers["Cookie"] = cookie
    return headers


def normalize_authorization(raw_value: str) -> str:
    """规范化 Authorization 配置。"""

    authorization = raw_value.strip()
    if authorization.lower().startswith("authorization:"):
        authorization = authorization.split(":", maxsplit=1)[1].strip()
    if authorization and not authorization.lower().startswith("bearer "):
        authorization = f"Bearer {authorization}"
    return authorization


def normalize_cookie(raw_value: str) -> str:
    """规范化 Cookie 配置。"""

    cookie = raw_value.strip()
    if cookie.lower().startswith("cookie:"):
        cookie = cookie.split(":", maxsplit=1)[1].strip()
    return cookie


def fetch_file_info(session: requests.Session, detailed_id: str) -> tuple[str, str]:
    """根据 detailed_id 获取 file_id 和 file_name。

    Args:
        session: 请求会话。
        detailed_id: 明细 ID。

    Returns:
        tuple[str, str]: 文件 ID 和文件名。
    """

    response = session.post(
        DETAIL_URL,
        json={"id": detailed_id},
        timeout=TIMEOUT,
    )
    response.raise_for_status()
    response_json = response.json()
    if str(response_json.get("code", "")).strip() != "00000":
        raise RuntimeError(
            f"详情接口返回异常: code={response_json.get('code')}, msg={response_json.get('msg')}"
        )
    data = response_json.get("data", {}) or {}
    file_id = str(data.get("file_id", "")).strip()
    file_name = str(data.get("file_name", "")).strip()
    return file_id, file_name


def download_file(
    session: requests.Session,
    file_id: str,
    output_path: Path,
    preferred_file_name: str,
) -> Path:
    """下载文件到本地。

    Args:
        session: 请求会话。
        file_id: 文件 ID。
        output_path: 输出文件路径。
        preferred_file_name: 详情接口返回的文件名。

    Returns:
        Path: 实际保存路径。
    """

    response = session.get(
        DOWNLOAD_URL,
        params={"fileKey": file_id},
        timeout=TIMEOUT,
        stream=True,
    )
    response.raise_for_status()
    actual_path = build_actual_output_path(
        file_id=file_id,
        output_path=output_path,
        response=response,
        preferred_file_name=preferred_file_name,
    )
    actual_path.parent.mkdir(parents=True, exist_ok=True)
    with actual_path.open("wb") as file_obj:
        for chunk in response.iter_content(chunk_size=8192):
            if chunk:
                file_obj.write(chunk)
    return actual_path


def build_actual_output_path(
    file_id: str,
    output_path: Path,
    response: requests.Response,
    preferred_file_name: str,
) -> Path:
    """根据响应内容生成最终保存路径。

    Args:
        file_id: 文件 ID。
        output_path: 默认输出路径。
        response: 下载响应对象。
        preferred_file_name: 详情接口返回的文件名。

    Returns:
        Path: 最终输出路径。
    """

    if preferred_file_name.strip():
        return output_path.with_name(sanitize_filename(preferred_file_name))
    filename = parse_filename_from_response(response=response)
    if not filename:
        return output_path.with_name(f"{file_id}.bin")
    return output_path.with_name(sanitize_filename(filename))


def parse_filename_from_response(response: requests.Response) -> str:
    """从响应头中提取文件名。

    Args:
        response: 下载响应对象。

    Returns:
        str: 文件名。
    """

    content_disposition = response.headers.get("Content-Disposition", "")
    if "filename=" not in content_disposition:
        return ""
    filename = content_disposition.split("filename=", maxsplit=1)[1].strip().strip('"')
    return filename


def sanitize_filename(filename: str) -> str:
    """清洗响应头中的文件名。"""

    normalized = filename.strip().replace("\\", "/")
    safe_name = Path(normalized).name
    if not safe_name:
        return "download.bin"
    return safe_name


def update_result(
    worksheet: Worksheet,
    row_index: int,
    file_id_column: int,
    file_name_column: int,
    status_column: int,
    path_column: int,
    error_column: int,
    file_id: str,
    file_name: str,
    status: str,
    download_path: str,
    error_message: str,
) -> None:
    """回写下载结果。"""

    worksheet.cell(row=row_index, column=file_id_column, value=file_id)
    worksheet.cell(row=row_index, column=file_name_column, value=file_name)
    worksheet.cell(row=row_index, column=status_column, value=status)
    worksheet.cell(row=row_index, column=path_column, value=download_path)
    worksheet.cell(row=row_index, column=error_column, value=error_message)


def get_cell_value(worksheet: Worksheet, row_index: int, column_name: str) -> str:
    """获取单元格字符串值。"""

    header_map = {
        str(worksheet.cell(row=1, column=column_index).value).strip(): column_index
        for column_index in range(1, worksheet.max_column + 1)
    }
    column_index = header_map[column_name]
    cell_value = worksheet.cell(row=row_index, column=column_index).value
    return "" if cell_value is None else str(cell_value).strip()


def build_output_path(file_id: str) -> Path:
    """构建输出路径。"""

    safe_file_id = sanitize_filename(file_id)
    if not safe_file_id:
        safe_file_id = "download.bin"
    return DOWNLOAD_DIR / safe_file_id


def main() -> None:
    """执行下载任务。"""

    if not AUTHORIZATION.strip() and not COOKIE.strip():
        raise ValueError("请先在脚本顶部配置 AUTHORIZATION 或 COOKIE。")

    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook[SHEET_NAME]
    file_id_column = get_or_create_column(worksheet, FILE_ID_COLUMN)
    file_name_column = get_or_create_column(worksheet, FILE_NAME_COLUMN)
    status_column = get_or_create_column(worksheet, DOWNLOAD_STATUS_COLUMN)
    path_column = get_or_create_column(worksheet, DOWNLOAD_PATH_COLUMN)
    error_column = get_or_create_column(worksheet, DOWNLOAD_ERROR_COLUMN)
    session = requests.Session()
    session.headers.update(build_headers())
    detail_cache: dict[str, tuple[str, str]] = {}
    download_cache: dict[str, Path] = {}
    for row_index in range(2, worksheet.max_row + 1):
        detailed_id = get_cell_value(worksheet, row_index, "detailed_id")
        if not detailed_id:
            update_result(
                worksheet=worksheet,
                row_index=row_index,
                file_id_column=file_id_column,
                file_name_column=file_name_column,
                status_column=status_column,
                path_column=path_column,
                error_column=error_column,
                file_id="",
                file_name="",
                status="skip",
                download_path="",
                error_message="detailed_id为空",
            )
            continue
        try:
            if detailed_id in detail_cache:
                file_id, file_name = detail_cache[detailed_id]
            else:
                file_id, file_name = fetch_file_info(session=session, detailed_id=detailed_id)
                detail_cache[detailed_id] = (file_id, file_name)
            if not file_id:
                update_result(
                    worksheet=worksheet,
                    row_index=row_index,
                    file_id_column=file_id_column,
                    file_name_column=file_name_column,
                    status_column=status_column,
                    path_column=path_column,
                    error_column=error_column,
                    file_id="",
                    file_name="",
                    status="failed",
                    download_path="",
                    error_message="未获取到file_id",
                )
                continue
            if file_id in download_cache:
                actual_path = download_cache[file_id]
                download_status = "success_cached"
            else:
                output_path = build_output_path(file_id=file_id)
                actual_path = download_file(
                    session=session,
                    file_id=file_id,
                    output_path=output_path,
                    preferred_file_name=file_name,
                )
                download_cache[file_id] = actual_path
                download_status = "success"
            update_result(
                worksheet=worksheet,
                row_index=row_index,
                file_id_column=file_id_column,
                file_name_column=file_name_column,
                status_column=status_column,
                path_column=path_column,
                error_column=error_column,
                file_id=file_id,
                file_name=file_name,
                status=download_status,
                download_path=str(actual_path),
                error_message="",
            )
            print(f"处理完成: detailed_id={detailed_id}, file_id={file_id}, file_name={file_name}, status={download_status}")
        except Exception as exc:
            update_result(
                worksheet=worksheet,
                row_index=row_index,
                file_id_column=file_id_column,
                file_name_column=file_name_column,
                status_column=status_column,
                path_column=path_column,
                error_column=error_column,
                file_id="",
                file_name="",
                status="failed",
                download_path="",
                error_message=str(exc),
            )
            print(f"下载失败: detailed_id={detailed_id}, error={exc}")
    workbook.save(EXCEL_PATH)


if __name__ == "__main__":
    main()
