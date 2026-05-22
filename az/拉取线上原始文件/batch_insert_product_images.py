from __future__ import annotations

import argparse
import json
import math
import random
import shutil
from pathlib import Path

from docx import Document
from docx.image.exceptions import UnrecognizedImageError
from docx.shared import Inches


DEFAULT_SOURCE_DIR = Path("/Users/layla.zhang/Downloads/推文文件/docx")
DEFAULT_IMAGE_ROOT_DIR = Path("/Users/layla.zhang/workspace/nullht-test/az/产品图片/logo")
DEFAULT_OUTPUT_DIR = Path("/Users/layla.zhang/测试用例/测试材料/az/产品图片原始推文材料")
SUPPORTED_DOCUMENT_EXTENSIONS = {".docx"}
SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
MIN_IMAGES_PER_FILE = 2
MAX_IMAGES_PER_FILE = 4
DEFAULT_SEED = 20260521
MANIFEST_NAME = "batch_insert_manifest.json"
OUTPUT_SUBDIR_NAME = "插入产品图片结果"
IMAGE_WIDTH_INCH = 2.2


def parse_args() -> argparse.Namespace:
    """解析命令行参数。

    Returns:
        argparse.Namespace: 参数对象。
    """

    parser = argparse.ArgumentParser(description="批量给 DOCX 插入单一产品分类图片")
    parser.add_argument(
        "--source-dir",
        default=str(DEFAULT_SOURCE_DIR),
        help="待处理 DOCX 目录",
    )
    parser.add_argument(
        "--image-root-dir",
        default=str(DEFAULT_IMAGE_ROOT_DIR),
        help="产品分类图片根目录",
    )
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR / OUTPUT_SUBDIR_NAME),
        help="输出目录",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=DEFAULT_SEED,
        help="随机种子，默认 20260521",
    )
    return parser.parse_args()


def list_documents(source_dir: Path) -> list[Path]:
    """收集待处理 DOCX 文件。

    Args:
        source_dir: 源目录。

    Returns:
        list[Path]: DOCX 文件列表。
    """

    return sorted(
        [
            file_path
            for file_path in source_dir.iterdir()
            if file_path.is_file()
            and file_path.suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS
            and not file_path.name.startswith(".")
        ]
    )


def list_category_images(image_root_dir: Path) -> dict[str, list[Path]]:
    """收集每个产品分类下的图片。

    Args:
        image_root_dir: 产品分类图片根目录。

    Returns:
        dict[str, list[Path]]: 分类到图片列表的映射。
    """

    category_images: dict[str, list[Path]] = {}
    for category_dir in sorted(image_root_dir.iterdir()):
        if not category_dir.is_dir() or category_dir.name.startswith("."):
            continue
        image_paths = sorted(
            [
                file_path
                for file_path in category_dir.iterdir()
                if file_path.is_file()
                and file_path.suffix.lower() in SUPPORTED_IMAGE_EXTENSIONS
                and not file_path.name.startswith(".")
            ]
        )
        if image_paths:
            category_images[category_dir.name] = image_paths
    return category_images


def ensure_minimum_image_count(image_paths: list[Path]) -> list[Path]:
    """确保单个文件至少插入 2 张图片。

    Args:
        image_paths: 原始图片列表。

    Returns:
        list[Path]: 补足后的图片列表。
    """

    if not image_paths:
        return []
    normalized_images = list(image_paths)
    while len(normalized_images) < MIN_IMAGES_PER_FILE:
        normalized_images.append(normalized_images[len(normalized_images) % len(image_paths)])
    return normalized_images


def build_assignment_plan(
    source_files: list[Path],
    category_images: dict[str, list[Path]],
) -> list[dict[str, object]]:
    """生成文件与图片的分配计划。

    Args:
        source_files: 原始 DOCX 文件列表。
        category_images: 产品分类图片映射。

    Returns:
        list[dict[str, object]]: 分配计划。
    """

    assignments: list[dict[str, object]] = []
    file_index = 0
    total_files = len(source_files)

    for category_name, image_paths in category_images.items():
        chunk_count = math.ceil(len(image_paths) / MAX_IMAGES_PER_FILE)
        for chunk_index in range(chunk_count):
            start_index = chunk_index * MAX_IMAGES_PER_FILE
            end_index = start_index + MAX_IMAGES_PER_FILE
            selected_images = image_paths[start_index:end_index]
            if len(selected_images) < MIN_IMAGES_PER_FILE and len(image_paths) >= MIN_IMAGES_PER_FILE:
                selected_images = image_paths[-MIN_IMAGES_PER_FILE:]
            selected_images = ensure_minimum_image_count(selected_images)

            source_path = source_files[file_index % total_files]
            copy_index = file_index // total_files
            assignments.append(
                {
                    "source_path": source_path,
                    "copy_index": copy_index,
                    "category_name": category_name,
                    "image_paths": selected_images,
                }
            )
            file_index += 1
    return assignments


def build_output_name(source_path: Path, copy_index: int, category_name: str) -> str:
    """构建输出文件名。

    Args:
        source_path: 原始文件路径。
        copy_index: 副本序号。
        category_name: 产品分类。

    Returns:
        str: 输出文件名。
    """

    suffix = "" if copy_index == 0 else f"_copy{copy_index + 1}"
    return f"{source_path.stem}__{category_name}{suffix}{source_path.suffix}"


def insert_images_into_document(
    source_path: Path,
    output_path: Path,
    category_name: str,
    image_paths: list[Path],
) -> list[dict[str, object]]:
    """给单个 DOCX 插入图片并保存。

    Args:
        source_path: 原始 DOCX 路径。
        output_path: 输出 DOCX 路径。
        category_name: 产品分类。
        image_paths: 待插入图片列表。

    Returns:
        list[dict[str, object]]: 实际插入记录。
    """

    document = Document(str(source_path))
    inserted_records: list[dict[str, object]] = []

    document.add_paragraph("")
    title_paragraph = document.add_paragraph()
    title_run = title_paragraph.add_run(f"产品图片分类：{category_name}")
    title_run.bold = True

    for image_path in image_paths:
        paragraph = document.add_paragraph()
        run = paragraph.add_run()
        try:
            run.add_picture(str(image_path), width=Inches(IMAGE_WIDTH_INCH))
        except UnrecognizedImageError:
            invalid_paragraph = document.add_paragraph(
                f"跳过不可识别图片：{image_path.name}"
            )
            inserted_records.append(
                {
                    "image_name": image_path.name,
                    "caption": invalid_paragraph.text,
                    "status": "skipped_unrecognized",
                }
            )
            continue
        caption_paragraph = document.add_paragraph(f"图片文件：{image_path.name}")
        inserted_records.append(
            {
                "image_name": image_path.name,
                "caption": caption_paragraph.text,
                "status": "inserted",
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return inserted_records


def write_manifest(output_dir: Path, manifest_rows: list[dict[str, object]]) -> Path:
    """写入处理结果清单。

    Args:
        output_dir: 输出目录。
        manifest_rows: 处理结果列表。

    Returns:
        Path: 清单文件路径。
    """

    manifest_path = output_dir / MANIFEST_NAME
    manifest_path.write_text(
        json.dumps(manifest_rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return manifest_path


def validate_inputs(source_files: list[Path], category_images: dict[str, list[Path]]) -> None:
    """校验输入。

    Args:
        source_files: 原始 DOCX 文件列表。
        category_images: 分类图片映射。
    """

    if not source_files:
        raise ValueError("source-dir 下未找到 docx 文件")
    if not category_images:
        raise ValueError("image-root-dir 下未找到分类图片")


def main() -> None:
    """执行批量插图。"""

    args = parse_args()
    random.seed(args.seed)
    source_dir = Path(args.source_dir).expanduser().resolve()
    image_root_dir = Path(args.image_root_dir).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    source_files = list_documents(source_dir)
    category_images = list_category_images(image_root_dir)
    validate_inputs(source_files, category_images)
    assignments = build_assignment_plan(source_files, category_images)

    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    manifest_rows: list[dict[str, object]] = []
    for assignment in assignments:
        source_path = assignment["source_path"]
        copy_index = int(assignment["copy_index"])
        category_name = str(assignment["category_name"])
        image_paths = list(assignment["image_paths"])
        output_name = build_output_name(
            source_path=source_path,
            copy_index=copy_index,
            category_name=category_name,
        )
        output_path = output_dir / output_name
        inserted_records = insert_images_into_document(
            source_path=source_path,
            output_path=output_path,
            category_name=category_name,
            image_paths=image_paths,
        )
        manifest_rows.append(
            {
                "source_file": str(source_path),
                "output_file": str(output_path),
                "category_name": category_name,
                "copy_index": copy_index,
                "image_names": [image_path.name for image_path in image_paths],
                "inserted_records": inserted_records,
            }
        )
        print(
            f"处理完成: output={output_path.name}, category={category_name}, "
            f"image_count={len(image_paths)}"
        )

    manifest_path = write_manifest(output_dir=output_dir, manifest_rows=manifest_rows)
    print(f"清单已写入: {manifest_path}")


if __name__ == "__main__":
    main()
