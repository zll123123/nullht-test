from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


EXCEL_PATH = Path("/Users/layla.zhang/workspace/nullht-test/az/拉取线上原始文件/结果统计base.xlsx")
SHEET_NAME = "明细"
ACTUAL_COUNT_COLUMN = "act_number"
PAGE_NO_COLUMN = "page_no"
EXPECT_NUMBER_COLUMN = "expect_number"
COMPARE_RESULT_COLUMN = "api_compare_result"
COMPARE_ERROR_COLUMN = "api_compare_error"


def find_column(worksheet: Worksheet, column_name: str) -> int:
    """查找表头列号。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。

    Returns:
        int: 对应列号。
    """

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    raise ValueError(f"Excel 缺少表头: {column_name}")


def find_optional_column(worksheet: Worksheet, column_name: str) -> int | None:
    """查找可选表头列号。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。

    Returns:
        int | None: 找到则返回列号，否则返回 None。
    """

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    return None


def get_or_create_column(worksheet: Worksheet, column_name: str) -> int:
    """获取或创建表头列。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。

    Returns:
        int: 对应列号。
    """

    column_index = find_optional_column(worksheet, column_name)
    if column_index is not None:
        return column_index
    new_column_index = worksheet.max_column + 1
    worksheet.cell(row=1, column=new_column_index, value=column_name)
    return new_column_index


def get_cell_value(worksheet: Worksheet, row_index: int, column_index: int) -> str:
    """读取单元格文本。

    Args:
        worksheet: Excel 工作表。
        row_index: 行号。
        column_index: 列号。

    Returns:
        str: 去空格后的文本值。
    """

    cell_value = worksheet.cell(row=row_index, column=column_index).value
    return "" if cell_value is None else str(cell_value).strip()


def normalize_page_no(raw_value: str) -> int | None:
    """规范化页码。

    Args:
        raw_value: 原始页码文本。

    Returns:
        int | None: 合法页码返回整数，否则返回 None。
    """

    if not raw_value:
        return None
    try:
        return int(float(raw_value))
    except ValueError:
        return None


def normalize_count(raw_value: str) -> int | None:
    """规范化数量。

    Args:
        raw_value: 原始数量文本。

    Returns:
        int | None: 合法数量返回整数，否则返回 None。
    """

    if not raw_value:
        return None
    try:
        return int(float(raw_value))
    except ValueError:
        return None


def write_compare_result(
    worksheet: Worksheet,
    row_index: int,
    compare_result_column: int,
    compare_error_column: int,
    compare_result: str,
    compare_error: str,
) -> None:
    """回写对比结果。

    Args:
        worksheet: Excel 工作表。
        row_index: 行号。
        compare_result_column: 结果列号。
        compare_error_column: 错误列号。
        compare_result: 对比结果。
        compare_error: 错误信息。
    """

    worksheet.cell(row=row_index, column=compare_result_column, value=compare_result)
    worksheet.cell(row=row_index, column=compare_error_column, value=compare_error)


def main() -> None:
    """仅基于 Excel 现有字段执行对比。"""

    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook[SHEET_NAME]
    page_no_column = find_column(worksheet, PAGE_NO_COLUMN)
    expect_number_column = find_column(worksheet, EXPECT_NUMBER_COLUMN)
    actual_count_column = find_column(worksheet, ACTUAL_COUNT_COLUMN)
    compare_result_column = get_or_create_column(worksheet, COMPARE_RESULT_COLUMN)
    compare_error_column = get_or_create_column(worksheet, COMPARE_ERROR_COLUMN)

    for row_index in range(2, worksheet.max_row + 1):
        page_no = normalize_page_no(get_cell_value(worksheet, row_index, page_no_column))
        expect_number = normalize_count(
            get_cell_value(worksheet, row_index, expect_number_column)
        )
        actual_count = normalize_count(
            get_cell_value(worksheet, row_index, actual_count_column)
        )

        invalid_fields: list[str] = []
        if page_no is None:
            invalid_fields.append("page_no")
        if expect_number is None:
            invalid_fields.append("expect_number")
        if actual_count is None:
            invalid_fields.append("act_number")

        if invalid_fields:
            write_compare_result(
                worksheet=worksheet,
                row_index=row_index,
                compare_result_column=compare_result_column,
                compare_error_column=compare_error_column,
                compare_result="blocked",
                compare_error=f"缺少或非法字段: {','.join(invalid_fields)}",
            )
            continue

        compare_result = "match" if expect_number == actual_count else "mismatch"
        write_compare_result(
            worksheet=worksheet,
            row_index=row_index,
            compare_result_column=compare_result_column,
            compare_error_column=compare_error_column,
            compare_result=compare_result,
            compare_error="",
        )

    workbook.save(EXCEL_PATH)


if __name__ == "__main__":
    main()
