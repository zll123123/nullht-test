from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import requests


EXCEL_PATH = Path("/Users/layla.zhang/workspace/nullht-test/az/拉取线上原始文件/结果统计base.xlsx")
SHEET_NAME = "明细"
SUMMARY_SHEET_NAME = "审计结果"
DETAIL_URL = "https://dev-api-v3-az-mlr.nullht.com/api/audit/management/detail"
TIMEOUT = 30
AUTHORIZATION = (
    "Bearer eyJ0eXAiOiJKV1QiLCJhbGciOiJIUzI1NiJ9.eyJsb2dpblR5cGUiOiJsb2dpbiIsImxvZ2luSWQiOiIwOGY4ZjU5NTQ4ZDQ0ODg4Yjk4ZTQ5OTkwYjUxMDQ5ZSIsImRldmljZVR5cGUiOiJERUYiLCJlZmYiOjE3NzkzMzg5MzMwNTUsInJuU3RyIjoib3I0NWhIeTZ5bEs2Z1NaRkFXUjFieUt6dDFQdXVadGUiLCJhel9vcGVuX2lkIjoieHh4eCJ9.AS3Jg60uKlYqzn7uDAq4PLCZWrgGNXAD6pQ70mVcSJs"
)
COOKIE = "acw_tc=65859a8117792557437945949ece6389d919e08c183cabe757a99998823922"

DETAIL_NEW_COLUMN = "detail_new"
TASK_ID_COLUMN = "task_id"
FILE_NAME_COLUMN = "file_name"
PAGE_NO_COLUMN = "page_no"
EXPECT_NUMBER_COLUMN = "expect_number"
ACTUAL_COUNT_COLUMN = "act_number"
COMPARE_RESULT_COLUMN = "api_compare_result"
COMPARE_ERROR_COLUMN = "api_compare_error"
RAW_API_DATA_COLUMN = "raw_api_data"
RAW_FINDINGS_COLUMN = "raw_product_image_findings"
SUMMARY_HEADERS = [
    "detail_new",
    "task_id",
    "file_name",
    "page_no",
    "expect_number",
    "act_number",
    "api_compare_result",
]
SUMMARY_METRIC_HEADERS = ["metric", "value"]
SUMMARY_METRIC_ORDER = [
    ("total_pages", "总页数"),
    ("match_pages", "一致通过页数"),
    ("mismatch_pages", "不一致页数"),
    ("tp", "TP"),
    ("fp", "FP"),
    ("fn", "FN"),
    ("recall", "召回率"),
    ("precision", "精确率"),
    ("f1_score", "F1 Score"),
]
DEFAULT_STEP = "all"
STEP_FETCH_RAW = "fetch_raw"
STEP_EXTRACT_FINDINGS = "extract_findings"
STEP_COMPARE = "compare"
STEP_SUMMARY = "summary"
VALID_STEPS = [
    DEFAULT_STEP,
    STEP_FETCH_RAW,
    STEP_EXTRACT_FINDINGS,
    STEP_COMPARE,
    STEP_SUMMARY,
]
SUMMARY_HEADER_ROW = len(SUMMARY_METRIC_ORDER) + 3
SUMMARY_DETAIL_START_ROW = SUMMARY_HEADER_ROW + 1


def parse_args() -> argparse.Namespace:
    """解析命令行参数。

    Returns:
        argparse.Namespace: 参数对象。
    """

    parser = argparse.ArgumentParser(description="产品图片数量对比工具")
    parser.add_argument(
        "--step",
        choices=VALID_STEPS,
        default=DEFAULT_STEP,
        help="执行步骤，可选: all/fetch_raw/extract_findings/compare/summary",
    )
    return parser.parse_args()


def build_headers() -> dict[str, str]:
    """构造请求头。

    Returns:
        dict[str, str]: 请求头字典。
    """

    return {
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh;q=0.9",
        "Authorization": AUTHORIZATION,
        "Connection": "keep-alive",
        "Content-Type": "application/json",
        "Cookie": COOKIE,
        "Origin": "https://dev-v3-az-mlr.nullht.com",
        "Referer": "https://dev-v3-az-mlr.nullht.com/",
        "Sec-Fetch-Dest": "empty",
        "Sec-Fetch-Mode": "cors",
        "Sec-Fetch-Site": "same-site",
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36"
        ),
        "sec-ch-ua": '"Chromium";v="148", "Google Chrome";v="148", "Not/A)Brand";v="99"',
        "sec-ch-ua-mobile": "?0",
        "sec-ch-ua-platform": '"macOS"',
    }


def find_column(worksheet: Worksheet, column_name: str) -> int:
    """查找表头列号。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。

    Returns:
        int: 列号。
    """

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    raise ValueError(f"Excel 缺少表头: {column_name}")


def get_or_create_column(worksheet: Worksheet, column_name: str) -> int:
    """获取或创建表头列。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。

    Returns:
        int: 列号。
    """

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    new_column_index = worksheet.max_column + 1
    worksheet.cell(row=1, column=new_column_index, value=column_name)
    return new_column_index


def find_optional_column(worksheet: Worksheet, column_name: str) -> int | None:
    """查找可选表头列号。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。

    Returns:
        int | None: 列号或 None。
    """

    for column_index in range(1, worksheet.max_column + 1):
        header_value = worksheet.cell(row=1, column=column_index).value
        if str(header_value).strip() == column_name:
            return column_index
    return None


def get_cell_value(worksheet: Worksheet, row_index: int, column_index: int) -> str:
    """读取单元格文本。

    Args:
        worksheet: Excel 工作表。
        row_index: 行号。
        column_index: 列号。

    Returns:
        str: 文本值。
    """

    cell_value = worksheet.cell(row=row_index, column=column_index).value
    return "" if cell_value is None else str(cell_value).strip()


def normalize_page_no(raw_value: str) -> int | None:
    """规范化页码。

    Args:
        raw_value: 原始页码。

    Returns:
        int | None: 页码整数或 None。
    """

    if not raw_value:
        return None
    try:
        return int(float(raw_value))
    except ValueError:
        return None


def normalize_count(raw_value: str) -> int:
    """规范化数量。

    Args:
        raw_value: 原始数量。

    Returns:
        int: 数量整数。
    """

    if not raw_value:
        return 0
    return int(float(raw_value))


def fetch_detail(session: requests.Session, detail_id: str) -> dict:
    """调用详情接口。

    Args:
        session: 请求会话。
        detail_id: 详情 ID。

    Returns:
        dict: 详情数据。
    """

    response = session.post(DETAIL_URL, json={"id": detail_id}, timeout=TIMEOUT)
    response.raise_for_status()
    response_json = response.json()
    if str(response_json.get("code", "")).strip() != "00000":
        raise RuntimeError(
            f"详情接口返回异常: code={response_json.get('code')}, msg={response_json.get('msg')}"
        )
    return response_json.get("data", {}) or {}


def is_product_image_finding(finding: dict) -> bool:
    """判断是否为产品图片审核点。

    Args:
        finding: 单个审核点数据。

    Returns:
        bool: 是否为产品图片审核点。
    """

    point_name = str(finding.get("point_name") or "").strip()
    return "产品图片" in point_name


def resolve_page_number(default_page_number: object, finding: dict) -> int | None:
    """解析错误所属页码。

    Args:
        default_page_number: 页级默认页码。
        finding: 单个审核点数据。

    Returns:
        int | None: 页码整数或 None。
    """

    positions = finding.get("position") or []
    if positions:
        page_number = positions[0].get("page_number")
        normalized_page_number = normalize_page_no("" if page_number is None else str(page_number))
        if normalized_page_number is not None:
            return normalized_page_number
    return normalize_page_no("" if default_page_number is None else str(default_page_number))


def build_raw_api_data(detail_data: dict) -> str:
    """构造仅包含产品图片审核点的原始数据字符串。

    Args:
        detail_data: 接口详情数据。

    Returns:
        str: JSON 字符串。
    """

    product_image_findings = extract_product_image_findings(detail_data)
    return json.dumps(product_image_findings, ensure_ascii=False)


def extract_product_image_findings(detail_data: dict) -> list[dict]:
    """提取产品图片相关审核点。

    Args:
        detail_data: 接口详情数据。

    Returns:
        list[dict]: 提取后的审核点列表。
    """

    raw_items: list[dict] = []
    for page_detail in detail_data.get("page_details") or []:
        default_page_number = page_detail.get("page_number")
        for finding in page_detail.get("findings") or []:
            if not is_product_image_finding(finding):
                continue
            raw_items.append(
                {
                    "page_number": resolve_page_number(default_page_number, finding),
                    "point_name": finding.get("point_name"),
                }
            )
    return raw_items


def build_page_count_map_from_findings(findings: list[dict]) -> dict[int, int]:
    """按页统计产品图片审核点数量。

    Args:
        findings: 产品图片审核点列表。

    Returns:
        dict[int, int]: 页码到数量映射。
    """

    page_count_map: dict[int, int] = {}
    for finding in findings:
        page_number = normalize_page_no(str(finding.get("page_number") or ""))
        if page_number is None:
            continue
        page_count_map[page_number] = page_count_map.get(page_number, 0) + 1
    return page_count_map


def resolve_detail_id(
    file_name: str,
    detail_new: str,
    last_file_name: str,
    last_detail_id: str,
) -> tuple[str, str, str]:
    """按连续 file_name 分组解析 detail_id。

    Args:
        file_name: 当前文件名。
        detail_new: 当前明细 ID。
        last_file_name: 上一行文件名。
        last_detail_id: 上一行明细 ID。

    Returns:
        tuple[str, str, str]: 解析后的明细 ID、最新文件名、最新明细 ID。
    """

    if detail_new:
        return detail_new, file_name, detail_new
    if file_name and file_name == last_file_name:
        return last_detail_id, last_file_name, last_detail_id
    return "", file_name, last_detail_id


def write_compare_result(
    worksheet: Worksheet,
    row_index: int,
    actual_count_column: int,
    compare_result_column: int,
    compare_error_column: int,
    actual_count: int,
    compare_result: str,
    compare_error: str,
) -> None:
    """回写对比结果。

    Args:
        worksheet: Excel 工作表。
        row_index: 行号。
        actual_count_column: 实际数量列号。
        compare_result_column: 对比结果列号。
        compare_error_column: 错误信息列号。
        actual_count: 实际数量。
        compare_result: 对比结果。
        compare_error: 错误信息。
    """

    worksheet.cell(row=row_index, column=actual_count_column, value=actual_count)
    worksheet.cell(row=row_index, column=compare_result_column, value=compare_result)
    worksheet.cell(row=row_index, column=compare_error_column, value=compare_error)


def write_text_cell(
    worksheet: Worksheet,
    row_index: int,
    column_index: int,
    text: str,
) -> None:
    """回写文本单元格。

    Args:
        worksheet: Excel 工作表。
        row_index: 行号。
        column_index: 列号。
        text: 文本值。
    """

    worksheet.cell(row=row_index, column=column_index, value=text)


def delete_column_if_exists(worksheet: Worksheet, column_name: str) -> None:
    """删除存在的列。

    Args:
        worksheet: Excel 工作表。
        column_name: 表头名称。
    """

    column_index = find_optional_column(worksheet, column_name)
    if column_index is None:
        return
    worksheet.delete_cols(column_index)


def ensure_compare_result_after_act_number(worksheet: Worksheet) -> tuple[int, int, int]:
    """确保结果列位于 act_number 后。

    Args:
        worksheet: Excel 工作表。

    Returns:
        tuple[int, int, int]: act_number、result、error 三列列号。
    """

    actual_count_column = find_column(worksheet, ACTUAL_COUNT_COLUMN)
    compare_result_column = find_optional_column(worksheet, COMPARE_RESULT_COLUMN)

    if compare_result_column is None:
        worksheet.insert_cols(actual_count_column + 1)
        worksheet.cell(row=1, column=actual_count_column + 1, value=COMPARE_RESULT_COLUMN)
    elif compare_result_column != actual_count_column + 1:
        worksheet.insert_cols(actual_count_column + 1)
        target_column = actual_count_column + 1
        for row_index in range(1, worksheet.max_row + 1):
            worksheet.cell(
                row=row_index,
                column=target_column,
                value=worksheet.cell(row=row_index, column=compare_result_column + 1).value,
            )
        worksheet.delete_cols(compare_result_column + 1)

    actual_count_column = find_column(worksheet, ACTUAL_COUNT_COLUMN)
    compare_result_column = find_column(worksheet, COMPARE_RESULT_COLUMN)
    compare_error_column = get_or_create_column(worksheet, COMPARE_ERROR_COLUMN)
    return actual_count_column, compare_result_column, compare_error_column


def reset_summary_sheet(worksheet: Worksheet) -> None:
    """重置审计结果 sheet 并写入表头。

    Args:
        worksheet: 审计结果工作表。
    """

    if worksheet.max_row > 0:
        worksheet.delete_rows(1, worksheet.max_row)
    for column_index, header in enumerate(SUMMARY_METRIC_HEADERS, start=1):
        worksheet.cell(row=1, column=column_index, value=header)
    for column_index, header in enumerate(SUMMARY_HEADERS, start=1):
        worksheet.cell(row=SUMMARY_HEADER_ROW, column=column_index, value=header)


def write_summary_metrics(worksheet: Worksheet, metrics: dict[str, int | float]) -> None:
    """写入汇总统计。

    Args:
        worksheet: 审计结果工作表。
        metrics: 统计指标。
    """

    for row_index, (metric_key, metric_name) in enumerate(SUMMARY_METRIC_ORDER, start=2):
        worksheet.cell(row=row_index, column=1, value=metric_name)
        worksheet.cell(row=row_index, column=2, value=metrics[metric_key])


def append_summary_row(
    worksheet: Worksheet,
    row_index: int,
    detail_id: str,
    task_id: str,
    file_name: str,
    page_no: int,
    expect_number: int,
    actual_count: int,
    compare_result: str,
) -> None:
    """追加 mismatch 汇总记录。

    Args:
        worksheet: 审计结果工作表。
        row_index: 行号。
        detail_id: 明细 ID。
        task_id: 任务 ID。
        file_name: 文件名。
        page_no: 页码。
        expect_number: 期望数量。
        actual_count: 实际数量。
        compare_result: 对比结果。
    """

    worksheet.cell(row=row_index, column=1, value=detail_id)
    worksheet.cell(row=row_index, column=2, value=task_id)
    worksheet.cell(row=row_index, column=3, value=file_name)
    worksheet.cell(row=row_index, column=4, value=page_no)
    worksheet.cell(row=row_index, column=5, value=expect_number)
    worksheet.cell(row=row_index, column=6, value=actual_count)
    worksheet.cell(row=row_index, column=7, value=compare_result)


def update_summary_metrics(
    metrics: dict[str, int],
    expect_number: int,
    actual_count: int,
    compare_result: str,
) -> None:
    """累计汇总统计。

    Args:
        metrics: 统计指标。
        expect_number: 期望数量。
        actual_count: 实际数量。
        compare_result: 对比结果。
    """

    metrics["total_pages"] += 1
    if compare_result == "match":
        metrics["match_pages"] += 1
    else:
        metrics["mismatch_pages"] += 1
    metrics["tp"] += min(expect_number, actual_count)
    metrics["fp"] += max(actual_count - expect_number, 0)
    metrics["fn"] += max(expect_number - actual_count, 0)


def safe_divide(numerator: int, denominator: int) -> float:
    """执行安全除法。

    Args:
        numerator: 分子。
        denominator: 分母。

    Returns:
        float: 除法结果，分母为 0 时返回 0。
    """

    if denominator == 0:
        return 0.0
    return numerator / denominator


def finalize_summary_metrics(metrics: dict[str, int | float]) -> None:
    """补充召回率、精确率与 F1 Score。

    Args:
        metrics: 统计指标。
    """

    tp = int(metrics["tp"])
    fp = int(metrics["fp"])
    fn = int(metrics["fn"])
    recall = safe_divide(tp, tp + fn)
    precision = safe_divide(tp, tp + fp)
    f1_score = safe_divide(2 * precision * recall, precision + recall)
    metrics["recall"] = recall
    metrics["precision"] = precision
    metrics["f1_score"] = f1_score


def load_raw_api_findings(raw_api_data: str) -> list[dict]:
    """解析 raw_api_data JSON。

    Args:
        raw_api_data: 原始接口数据 JSON 字符串。

    Returns:
        list[dict]: 产品图片审核点列表。
    """

    if not raw_api_data:
        return []
    try:
        loaded_data = json.loads(raw_api_data)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "raw_api_data 解析失败，可能是 Excel 单元格内容超长被截断，请先重新执行 fetch_raw"
        ) from exc
    if not isinstance(loaded_data, list):
        raise ValueError("raw_api_data 不是合法产品图片审核点列表")
    return loaded_data


def load_raw_findings(raw_findings: str) -> list[dict]:
    """解析产品图片审核点 JSON。

    Args:
        raw_findings: 产品图片审核点 JSON 字符串。

    Returns:
        list[dict]: 审核点列表。
    """

    if not raw_findings:
        return []
    loaded_data = json.loads(raw_findings)
    if not isinstance(loaded_data, list):
        raise ValueError("raw_product_image_findings 不是合法列表")
    return loaded_data


def get_sheet_pair() -> tuple[object, Worksheet, Worksheet]:
    """加载工作簿与工作表。

    Returns:
        tuple[object, Worksheet, Worksheet]: 工作簿、明细表、汇总表。
    """

    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook[SHEET_NAME]
    summary_worksheet = workbook[SUMMARY_SHEET_NAME]
    return workbook, worksheet, summary_worksheet


def build_session() -> requests.Session:
    """构造请求会话。

    Returns:
        requests.Session: 请求会话。
    """

    session = requests.Session()
    session.headers.update(build_headers())
    return session


def step_fetch_raw_data() -> None:
    """步骤一：拉取接口原始数据并回写 raw_api_data。"""

    workbook, worksheet, _ = get_sheet_pair()
    detail_new_column = find_column(worksheet, DETAIL_NEW_COLUMN)
    file_name_column = find_column(worksheet, FILE_NAME_COLUMN)
    raw_api_data_column = get_or_create_column(worksheet, RAW_API_DATA_COLUMN)
    session = build_session()
    detail_data_cache: dict[str, str] = {}
    written_raw_detail_ids: set[str] = set()
    last_file_name = ""
    last_detail_id = ""

    for row_index in range(2, worksheet.max_row + 1):
        file_name = get_cell_value(worksheet, row_index, file_name_column)
        detail_new = get_cell_value(worksheet, row_index, detail_new_column)
        resolved_detail_id, last_file_name, last_detail_id = resolve_detail_id(
            file_name=file_name,
            detail_new=detail_new,
            last_file_name=last_file_name,
            last_detail_id=last_detail_id,
        )
        if not resolved_detail_id or resolved_detail_id in written_raw_detail_ids:
            continue
        detail_data_json = detail_data_cache.get(resolved_detail_id)
        if detail_data_json is None:
            detail_data = fetch_detail(session=session, detail_id=resolved_detail_id)
            detail_data_json = build_raw_api_data(detail_data)
            detail_data_cache[resolved_detail_id] = detail_data_json
        write_text_cell(
            worksheet=worksheet,
            row_index=row_index,
            column_index=raw_api_data_column,
            text=detail_data_json,
        )
        written_raw_detail_ids.add(resolved_detail_id)

    workbook.save(EXCEL_PATH)


def step_extract_findings() -> None:
    """步骤二：从 raw_api_data 中标准化提取产品图片审核点。"""

    workbook, worksheet, _ = get_sheet_pair()
    raw_api_data_column = find_column(worksheet, RAW_API_DATA_COLUMN)
    raw_findings_column = get_or_create_column(worksheet, RAW_FINDINGS_COLUMN)

    for row_index in range(2, worksheet.max_row + 1):
        raw_api_data = get_cell_value(worksheet, row_index, raw_api_data_column)
        if not raw_api_data:
            write_text_cell(worksheet, row_index, raw_findings_column, "")
            continue
        product_image_findings = load_raw_api_findings(raw_api_data)
        write_text_cell(
            worksheet=worksheet,
            row_index=row_index,
            column_index=raw_findings_column,
            text=json.dumps(product_image_findings, ensure_ascii=False),
        )

    workbook.save(EXCEL_PATH)


def build_findings_cache(worksheet: Worksheet) -> dict[str, dict[int, int]]:
    """基于 raw_api_data 构建 detail_id 到页数量映射缓存。

    Args:
        worksheet: 明细工作表。

    Returns:
        dict[str, dict[int, int]]: 缓存映射。
    """

    detail_new_column = find_column(worksheet, DETAIL_NEW_COLUMN)
    file_name_column = find_column(worksheet, FILE_NAME_COLUMN)
    raw_api_data_column = find_column(worksheet, RAW_API_DATA_COLUMN)
    findings_cache: dict[str, dict[int, int]] = {}
    last_file_name = ""
    last_detail_id = ""

    for row_index in range(2, worksheet.max_row + 1):
        file_name = get_cell_value(worksheet, row_index, file_name_column)
        detail_new = get_cell_value(worksheet, row_index, detail_new_column)
        resolved_detail_id, last_file_name, last_detail_id = resolve_detail_id(
            file_name=file_name,
            detail_new=detail_new,
            last_file_name=last_file_name,
            last_detail_id=last_detail_id,
        )
        raw_api_data = get_cell_value(worksheet, row_index, raw_api_data_column)
        if not resolved_detail_id or not raw_api_data or resolved_detail_id in findings_cache:
            continue
        findings_cache[resolved_detail_id] = build_page_count_map_from_findings(
            load_raw_api_findings(raw_api_data)
        )
    return findings_cache


def step_compare_counts() -> None:
    """步骤三：基于提取结果按页对比 expect_number 与 act_number。"""

    workbook, worksheet, _ = get_sheet_pair()
    delete_column_if_exists(worksheet, "api_product_image_count")
    delete_column_if_exists(worksheet, "api_compare_diff")
    detail_new_column = find_column(worksheet, DETAIL_NEW_COLUMN)
    file_name_column = find_column(worksheet, FILE_NAME_COLUMN)
    page_no_column = find_column(worksheet, PAGE_NO_COLUMN)
    expect_number_column = find_column(worksheet, EXPECT_NUMBER_COLUMN)
    actual_count_column, compare_result_column, compare_error_column = (
        ensure_compare_result_after_act_number(worksheet)
    )
    findings_cache = build_findings_cache(worksheet)
    last_file_name = ""
    last_detail_id = ""

    for row_index in range(2, worksheet.max_row + 1):
        file_name = get_cell_value(worksheet, row_index, file_name_column)
        detail_new = get_cell_value(worksheet, row_index, detail_new_column)
        resolved_detail_id, last_file_name, last_detail_id = resolve_detail_id(
            file_name=file_name,
            detail_new=detail_new,
            last_file_name=last_file_name,
            last_detail_id=last_detail_id,
        )
        page_no = normalize_page_no(get_cell_value(worksheet, row_index, page_no_column))
        expect_number = normalize_count(get_cell_value(worksheet, row_index, expect_number_column))

        if not resolved_detail_id or page_no is None:
            missing_fields: list[str] = []
            if not resolved_detail_id:
                missing_fields.append("detail_new")
            if page_no is None:
                missing_fields.append("page_no")
            write_compare_result(
                worksheet=worksheet,
                row_index=row_index,
                actual_count_column=actual_count_column,
                compare_result_column=compare_result_column,
                compare_error_column=compare_error_column,
                actual_count=0,
                compare_result="blocked",
                compare_error=f"缺少{','.join(missing_fields)}",
            )
            continue

        page_count_map = findings_cache.get(resolved_detail_id, {})
        actual_count = page_count_map.get(page_no, 0)
        compare_result = "match" if actual_count == expect_number else "mismatch"
        write_compare_result(
            worksheet=worksheet,
            row_index=row_index,
            actual_count_column=actual_count_column,
            compare_result_column=compare_result_column,
            compare_error_column=compare_error_column,
            actual_count=actual_count,
            compare_result=compare_result,
            compare_error="",
        )

    workbook.save(EXCEL_PATH)


def step_build_summary() -> None:
    """步骤四：汇总最终对比结果。"""

    workbook, worksheet, summary_worksheet = get_sheet_pair()
    detail_new_column = find_column(worksheet, DETAIL_NEW_COLUMN)
    task_id_column = find_column(worksheet, TASK_ID_COLUMN)
    file_name_column = find_column(worksheet, FILE_NAME_COLUMN)
    page_no_column = find_column(worksheet, PAGE_NO_COLUMN)
    expect_number_column = find_column(worksheet, EXPECT_NUMBER_COLUMN)
    actual_count_column = find_column(worksheet, ACTUAL_COUNT_COLUMN)
    compare_result_column = find_column(worksheet, COMPARE_RESULT_COLUMN)
    reset_summary_sheet(summary_worksheet)

    summary_row_index = SUMMARY_DETAIL_START_ROW
    summary_metrics: dict[str, int | float] = {
        "total_pages": 0,
        "match_pages": 0,
        "mismatch_pages": 0,
        "tp": 0,
        "fp": 0,
        "fn": 0,
        "recall": 0.0,
        "precision": 0.0,
        "f1_score": 0.0,
    }

    for row_index in range(2, worksheet.max_row + 1):
        compare_result = get_cell_value(worksheet, row_index, compare_result_column)
        if compare_result not in {"match", "mismatch"}:
            continue
        detail_id = get_cell_value(worksheet, row_index, detail_new_column)
        task_id = get_cell_value(worksheet, row_index, task_id_column)
        file_name = get_cell_value(worksheet, row_index, file_name_column)
        page_no = normalize_page_no(get_cell_value(worksheet, row_index, page_no_column))
        expect_number = normalize_count(get_cell_value(worksheet, row_index, expect_number_column))
        actual_count = normalize_count(get_cell_value(worksheet, row_index, actual_count_column))
        if page_no is None:
            continue
        update_summary_metrics(
            metrics=summary_metrics,
            expect_number=expect_number,
            actual_count=actual_count,
            compare_result=compare_result,
        )
        if compare_result == "mismatch":
            append_summary_row(
                worksheet=summary_worksheet,
                row_index=summary_row_index,
                detail_id=detail_id,
                task_id=task_id,
                file_name=file_name,
                page_no=page_no,
                expect_number=expect_number,
                actual_count=actual_count,
                compare_result=compare_result,
            )
            summary_row_index += 1

    finalize_summary_metrics(summary_metrics)
    write_summary_metrics(summary_worksheet, summary_metrics)
    workbook.save(EXCEL_PATH)


def run_all_steps() -> None:
    """执行全部步骤。"""

    step_fetch_raw_data()
    step_extract_findings()
    step_compare_counts()
    step_build_summary()


def main() -> None:
    """执行入口。"""

    args = parse_args()
    if args.step == DEFAULT_STEP:
        run_all_steps()
        return
    if args.step == STEP_FETCH_RAW:
        step_fetch_raw_data()
        return
    if args.step == STEP_EXTRACT_FINDINGS:
        step_extract_findings()
        return
    if args.step == STEP_COMPARE:
        step_compare_counts()
        return
    step_build_summary()


if __name__ == "__main__":
    main()
