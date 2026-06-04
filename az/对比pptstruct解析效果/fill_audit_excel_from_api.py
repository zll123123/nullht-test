#!/usr/bin/env python3
"""根据审核管理 API 回填 Excel 中的任务结果。

脚本作用：
1. 读取 `ppt原文抽取结果_获取pptstruct对比.xlsx` 中的 `文件名称` 列，去重后作为查询种子。
2. 调用 `/api/audit/management/list`，按文件名查询最新审核记录。
3. 再调用 `/api/audit/management/detail` 获取详情，抽取：
   - `taskid`
   - `detailId`
   - 每页的 `报错审核点`
   - 每页的 `审核点错误摘要`
4. 按 `文件名称 + 页码` 回填到 Excel 的现有行中。
5. 同名文件的所有行会统一写入相同的 `taskid` 和 `detailId`。

当前执行模式：
1. 默认模式：处理 active sheet。
   用法：`python fill_audit_excel_from_api.py`
2. 指定工作表模式：只切换目标 sheet，业务逻辑不变。
   用法：`python fill_audit_excel_from_api.py --sheet-name <sheet_name>`

补充说明：
1. 这是固定流程脚本，没有类似 error-page/full-file 的业务模式切换。
2. 脚本可直接通过 `python fill_audit_excel_from_api.py` 执行。
"""

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import requests
from loguru import logger
from openpyxl import load_workbook


DEFAULT_BASE_URL = "https://dev-az-ai-mlr-api.nullht.com"
DEFAULT_LIST_PATH = "/api/audit/management/list"
DEFAULT_DETAIL_PATH = "/api/audit/management/detail"
DEFAULT_EXCEL_PATH = Path(__file__).resolve().with_name("ppt原文抽取结果_获取pptstruct对比.xlsx")
DEFAULT_LOG_PATH = Path(__file__).resolve().with_name("fill_audit_excel_from_api.log")
DEFAULT_BEARER_TOKEN = "Bearer eyJ0eXAiOiJKV1QiLCJhbGciOiJIUzI1NiJ9.eyJsb2dpblR5cGUiOiJsb2dpbiIsImxvZ2luSWQiOiIyZmY5ZjY0ZmYwYmZmY2I0MzhkNmQ0MWViZDA4ZTQ5MyIsImRldmljZVR5cGUiOiJERUYiLCJlZmYiOjE3ODA0ODk0OTMzMzcsInJuU3RyIjoieWUyejI0a1laU0dpTUNkaGZ0SGFqNU5SaDNFWlVkU0sifQ.1z-N7ll1RYoB2SuUbwGI09_t-MBitUBcrX6nNN3GADE"
DEFAULT_COOKIE = ""
TIMEOUT_SECONDS = 60

FILE_NAME_HEADER = "文件名称"
TASK_ID_HEADER = "taskid"
DETAIL_ID_HEADER = "detailId"
PAGE_NUMBER_HEADER = "页码"
POINT_NAME_HEADER = "报错审核点"
REASON_HEADER = "审核点错误摘要"

COMMON_HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Connection": "keep-alive",
    "Origin": "https://dev-v4-az-mlr.nullht.com",
    "Referer": "https://dev-v4-az-mlr.nullht.com",
    "Sec-Fetch-Dest": "empty",
    "Sec-Fetch-Mode": "cors",
    "Sec-Fetch-Site": "same-site",
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0.0.0 Safari/537.36"
    ),
    "sec-ch-ua": '"Google Chrome";v="147", "Not.A/Brand";v="8", "Chromium";v="147"',
    "sec-ch-ua-mobile": "?0",
    "sec-ch-ua-platform": '"macOS"',
}

RUNTIME_CONFIG: Dict[str, Any] = {}


def parse_args() -> argparse.Namespace:
    """解析命令行参数。

    Returns:
        argparse.Namespace: 启动参数。
    """
    parser = argparse.ArgumentParser(description="根据 Excel 中的文件名称获取审核结果，并按页回填到现有行。")
    parser.add_argument("--sheet-name", help="可选，指定工作表名称；默认使用 active sheet")
    return parser.parse_args()


def configure_logger() -> None:
    """初始化日志输出。"""
    logger.remove()
    logger.add(sys.stdout, level="INFO", format="{message}")
    logger.add(DEFAULT_LOG_PATH, level="INFO", format="{message}", encoding="utf-8", mode="w")


def init_runtime_config(sheet_name: Optional[str] = None) -> None:
    """初始化运行时配置。

    Args:
        sheet_name: 工作表名称。
    """
    RUNTIME_CONFIG.update(
        {
            "bearer_token": DEFAULT_BEARER_TOKEN,
            "cookie": DEFAULT_COOKIE,
            "base_url": DEFAULT_BASE_URL,
            "sheet_name": sheet_name,
        }
    )


def build_headers(token: str, cookie: str) -> Dict[str, str]:
    """构造请求头。

    Args:
        token: Bearer token。
        cookie: Cookie。

    Returns:
        Dict[str, str]: 请求头。
    """
    authorization = token.strip()
    if not authorization.lower().startswith("bearer "):
        authorization = f"Bearer {authorization}"
    return {
        **COMMON_HEADERS,
        "Authorization": authorization,
        "Content-Type": "application/json",
        "Cookie": cookie,
    }


def normalize_header(value: Any) -> str:
    """标准化表头文本。

    Args:
        value: 表头值。

    Returns:
        str: 标准化后的结果。
    """
    if value is None:
        return ""
    return str(value).strip()


def find_header_map(sheet) -> Dict[str, int]:
    """读取表头映射。

    Args:
        sheet: 工作表。

    Returns:
        Dict[str, int]: 表头到列号映射。
    """
    header_map: Dict[str, int] = {}
    for column_index in range(1, sheet.max_column + 1):
        header_map[normalize_header(sheet.cell(1, column_index).value)] = column_index
    return header_map


def find_required_columns(sheet) -> Dict[str, int]:
    """获取必需列。

    Args:
        sheet: 工作表。

    Returns:
        Dict[str, int]: 字段到列号映射。
    """
    header_map = find_header_map(sheet)
    required_headers = {
        "file_name": FILE_NAME_HEADER,
        "task_id": TASK_ID_HEADER,
        "detail_id": DETAIL_ID_HEADER,
        "page_number": PAGE_NUMBER_HEADER,
        "point_name": POINT_NAME_HEADER,
        "reason": REASON_HEADER,
    }
    columns: Dict[str, int] = {}
    for key, header in required_headers.items():
        column_index = header_map.get(header)
        if not column_index:
            raise AssertionError(f"Excel 缺少表头: {header}")
        columns[key] = column_index
    return columns


def normalize_page_number(value: Any) -> Optional[str]:
    """标准化页码。

    Args:
        value: 原始页码。

    Returns:
        Optional[str]: 标准化结果。
    """
    if value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def load_seed_files(sheet, column_map: Dict[str, int]) -> List[str]:
    """读取待处理文件名列表。

    Args:
        sheet: 工作表。
        column_map: 列映射。

    Returns:
        List[str]: 去重后的文件名列表。
    """
    file_names: List[str] = []
    seen = set()
    for row_index in range(2, sheet.max_row + 1):
        value = sheet.cell(row_index, column_map["file_name"]).value
        if not value:
            continue
        file_name = str(value).strip()
        if not file_name or file_name in seen:
            continue
        file_names.append(file_name)
        seen.add(file_name)
    return file_names


def build_row_index(sheet, column_map: Dict[str, int]) -> Dict[Tuple[str, str], List[int]]:
    """建立按文件名和页码定位行的索引。

    Args:
        sheet: 工作表。
        column_map: 列映射。

    Returns:
        Dict[Tuple[str, str], List[int]]: 索引结果。
    """
    index: Dict[Tuple[str, str], List[int]] = {}
    for row_index in range(2, sheet.max_row + 1):
        file_name_value = sheet.cell(row_index, column_map["file_name"]).value
        page_number_value = sheet.cell(row_index, column_map["page_number"]).value
        if not file_name_value:
            continue
        page_number = normalize_page_number(page_number_value)
        if not page_number:
            continue
        key = (str(file_name_value).strip(), page_number)
        index.setdefault(key, []).append(row_index)
    return index


def build_file_row_index(sheet, column_map: Dict[str, int]) -> Dict[str, List[int]]:
    """建立按文件名定位所有行的索引。

    Args:
        sheet: 工作表。
        column_map: 列映射。

    Returns:
        Dict[str, List[int]]: 文件名到行号列表的映射。
    """
    index: Dict[str, List[int]] = {}
    for row_index in range(2, sheet.max_row + 1):
        file_name_value = sheet.cell(row_index, column_map["file_name"]).value
        if not file_name_value:
            continue
        file_name = str(file_name_value).strip()
        if not file_name:
            continue
        index.setdefault(file_name, []).append(row_index)
    return index


def mask_token(token: str) -> str:
    """脱敏 token。

    Args:
        token: 原始 token。

    Returns:
        str: 脱敏结果。
    """
    if len(token) <= 12:
        return "***"
    return f"{token[:6]}...{token[-6:]}"


def log_request(name: str, url: str, payload: Dict[str, Any], headers: Dict[str, str]) -> None:
    """打印请求日志。

    Args:
        name: 请求名。
        url: URL。
        payload: 请求体。
        headers: 请求头。
    """
    safe_headers = dict(headers)
    safe_headers["Authorization"] = f"Bearer {mask_token(headers['Authorization'].replace('Bearer ', '', 1))}"
    logger.info(f"===== {name} 请求开始 =====")
    logger.info(f"{name} URL: {url}")
    logger.info(f"{name} Headers: {json.dumps(safe_headers, ensure_ascii=False)}")
    logger.info(f"{name} Payload: {json.dumps(payload, ensure_ascii=False)}")


def post_json(session: requests.Session, url: str, payload: Dict[str, Any], headers: Dict[str, str], name: str) -> Dict[str, Any]:
    """发送 POST JSON 请求。

    Args:
        session: 请求会话。
        url: URL。
        payload: 请求体。
        headers: 请求头。
        name: 请求名。

    Returns:
        Dict[str, Any]: 响应 JSON。
    """
    log_request(name, url, payload, headers)
    response = session.post(url, json=payload, headers=headers, timeout=TIMEOUT_SECONDS)
    logger.info(f"{name} 响应状态码: {response.status_code}")
    logger.info(f"{name} 响应内容: {response.text}")
    logger.info(f"===== {name} 请求结束 =====")
    response.raise_for_status()
    return response.json()


def fetch_list_row(session: requests.Session, base_url: str, file_name: str, headers: Dict[str, str]) -> Optional[Dict[str, Any]]:
    """按文件名获取最新审核记录。

    Args:
        session: 请求会话。
        base_url: 域名。
        file_name: 文件名。
        headers: 请求头。

    Returns:
        Optional[Dict[str, Any]]: 列表首条记录。
    """
    response = post_json(
        session,
        f"{base_url}{DEFAULT_LIST_PATH}",
        {
            "page_num": 1,
            "page_size": 10,
            "status": [],
            "file_name": file_name,
            "sorts": ["CREATED_TIME_DESC"],
        },
        headers,
        "LIST",
    )
    if response.get("code") != "00000":
        raise RuntimeError(f"LIST 返回异常: {response.get('code')} {response.get('msg')}")
    return ((response.get("data") or {}).get("rows") or [None])[0]


def fetch_detail(session: requests.Session, base_url: str, row_id: str, headers: Dict[str, str]) -> Dict[str, Any]:
    """获取审核详情。

    Args:
        session: 请求会话。
        base_url: 域名。
        row_id: 详情 ID。
        headers: 请求头。

    Returns:
        Dict[str, Any]: 详情 JSON。
    """
    response = post_json(
        session,
        f"{base_url}{DEFAULT_DETAIL_PATH}",
        {"id": row_id},
        headers,
        "DETAIL",
    )
    if response.get("code") != "00000":
        raise RuntimeError(f"DETAIL 返回异常: {response.get('code')} {response.get('msg')}")
    return response.get("data") or {}


def collect_page_findings(detail: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
    """按页聚合审核点和错误摘要。

    Args:
        detail: 详情 JSON。

    Returns:
        Dict[str, Dict[str, str]]: 页码到聚合结果的映射。
    """
    result: Dict[str, Dict[str, List[str]]] = {}
    for page_detail in detail.get("page_details") or []:
        default_page_number = normalize_page_number(page_detail.get("page_number"))
        for finding in page_detail.get("findings") or []:
            page_number = default_page_number
            positions = finding.get("position") or []
            if positions and positions[0].get("page_number") is not None:
                page_number = normalize_page_number(positions[0].get("page_number"))
            if not page_number:
                continue
            page_result = result.setdefault(page_number, {"point_names": [], "reasons": []})
            point_name = str(finding.get("point_name") or "").strip()
            reason = str(finding.get("reason") or ((finding.get("extension") or {}).get("llm_reason")) or "").strip()
            if point_name and point_name not in page_result["point_names"]:
                page_result["point_names"].append(point_name)
            if reason:
                page_result["reasons"].append(reason)

    formatted: Dict[str, Dict[str, str]] = {}
    for page_number, values in result.items():
        formatted[page_number] = {
            "point_name": "/".join(values["point_names"]),
            "reason": "&&".join(values["reasons"]),
        }
    return formatted


def write_row(sheet, row_index: int, column_map: Dict[str, int], task_id: str, detail_id: str, point_name: str, reason: str) -> None:
    """更新单行结果。

    Args:
        sheet: 工作表。
        row_index: 行号。
        column_map: 列映射。
        task_id: 任务 ID。
        detail_id: 详情 ID。
        point_name: 审核点。
        reason: 错误摘要。
    """
    sheet.cell(row_index, column_map["task_id"]).value = task_id
    sheet.cell(row_index, column_map["detail_id"]).value = detail_id
    sheet.cell(row_index, column_map["point_name"]).value = point_name
    sheet.cell(row_index, column_map["reason"]).value = reason


def run_fill_audit_excel_from_api() -> None:
    """执行 API 回填主流程。"""
    if not RUNTIME_CONFIG:
        configure_logger()
        init_runtime_config()
    assert DEFAULT_BEARER_TOKEN.strip(), "请先在代码常量 DEFAULT_BEARER_TOKEN 中配置 token。"

    excel_path = DEFAULT_EXCEL_PATH.expanduser().resolve()
    workbook = load_workbook(excel_path)
    sheet = workbook[RUNTIME_CONFIG["sheet_name"]] if RUNTIME_CONFIG["sheet_name"] else workbook.active
    column_map = find_required_columns(sheet)
    seed_files = load_seed_files(sheet, column_map)
    row_index_map = build_row_index(sheet, column_map)
    file_row_index_map = build_file_row_index(sheet, column_map)
    headers = build_headers(RUNTIME_CONFIG["bearer_token"], RUNTIME_CONFIG["cookie"])
    session = requests.Session()

    logger.info(f"Excel 路径: {excel_path}")
    logger.info(f"日志路径: {DEFAULT_LOG_PATH}")
    logger.info(f"待处理文件数: {len(seed_files)}")

    updated_rows = 0
    for file_name in seed_files:
        try:
            list_row = fetch_list_row(session, RUNTIME_CONFIG["base_url"], file_name, headers)
            if not list_row:
                logger.info(f"未查询到审核记录: {file_name}")
                continue

            detail_id = str(list_row.get("id") or "").strip()
            detail = fetch_detail(session, RUNTIME_CONFIG["base_url"], detail_id, headers)
            task_id = str(detail.get("task_id") or list_row.get("task_id") or "").strip()
            page_findings = collect_page_findings(detail)

            for row_index in file_row_index_map.get(file_name, []):
                sheet.cell(row_index, column_map["task_id"]).value = task_id
                sheet.cell(row_index, column_map["detail_id"]).value = detail_id

            for page_number, finding in page_findings.items():
                row_indexes = row_index_map.get((file_name, page_number), [])
                for row_index in row_indexes:
                    write_row(
                        sheet,
                        row_index,
                        column_map,
                        task_id,
                        detail_id,
                        finding["point_name"],
                        finding["reason"],
                    )
                    updated_rows += 1
        except Exception as exc:
            logger.info(f"处理失败: file={file_name} error={exc}")

    workbook.save(excel_path)
    logger.info(f"Excel 已更新: {excel_path}")
    logger.info(f"写入记录数: {updated_rows}")


def main() -> int:
    """脚本主入口。

    Returns:
        int: 退出码。
    """
    args = parse_args()
    configure_logger()
    init_runtime_config(args.sheet_name)
    run_fill_audit_excel_from_api()
    return 0


if __name__ == "__main__":
    sys.exit(main())
