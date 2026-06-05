#!/usr/bin/env python3
"""根据日志中的 `getPptStruct` 结果，回填 Excel 里的日志字段。

脚本作用：
1. 读取 `ppt原文抽取结果_获取pptstruct对比.xlsx`。
2. 从 Excel 中读取：
   - `taskid`
   - `页码`
3. 扫描日志，提取两类与页码相关的日志数据：
   - `[任务编号][getPptStruct] 第 N 页解析成功,解析结果为：{...}`
   - `[任务编号][getPptStruct] 获取文件解析结果, page: {...}`
4. 按 `taskid + 页码` 回填以下列：
   - `日志中的pptstruct`
   - `日志paddleocr识别结果`
   - `日志中的图片解析内容`
   - `日志中的图表解析内容`

提取逻辑：
1. `日志中的pptstruct`
   - 直接使用 `第 N 页解析成功,解析结果为：{...}` 的单页结果。
2. `日志paddleocr识别结果`
   - 使用同任务同页最近一次 `获取文件解析结果, page: {...}` 中的 `content_blocks`。
   - 这部分就是提示词里 `1. OCR识别结果：` 对应的主体内容。
3. `日志中的图片解析内容`
   - 使用同任务同页最近一次 `获取文件解析结果, page: {...}` 中的 `image_results`。
4. `日志中的图表解析内容`
   - 使用同任务同页最近一次 `获取文件解析结果, page: {...}` 中的 `chart_results`。

执行模式：
1. `error-page`
   - 默认模式。
   - 只回填当前行页码对应的单页结果。
2. `full-file`
   - 将该任务下所有页的 `pptStruct` 聚合成整份 JSON，写回当前行的 `日志中的pptstruct`。
   - 另外三个日志字段仍按当前行的 `页码` 回填单页内容。

可选参数：
1. `--sheet-name <sheet_name>`
   - 指定工作表，默认使用 active sheet。
2. `--mode error-page`
   - 只提取当前页。
3. `--mode full-file`
   - 提取整份文件所有页的 `pptStruct`。
4. `--skip-non-empty`
   - 如果 `日志中的pptstruct` 已有内容，则跳过当前行不覆盖。
"""

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


DEFAULT_EXCEL_PATH = Path(__file__).resolve().with_name("ppt原文抽取结果_获取pptstruct对比.xlsx")
DEFAULT_LOG_SOURCE = Path(__file__).resolve().parent

TASK_ID_HEADER = "taskid"
PAGE_NUMBER_HEADER = "页码"
PPT_STRUCT_HEADER = "日志中的pptstruct"
PADDLE_OCR_HEADER = "日志paddleocr识别结果"
IMAGE_PARSE_HEADER = "日志中的图片解析内容"
CHART_PARSE_HEADER = "日志中的图表解析内容"

PAGE_PPT_STRUCT_PATTERN = re.compile(
    r"\[(?P<task_id>[A-Za-z0-9]+)\]\[getPptStruct\]\s*第\s*(?P<page_number>\d+)\s*页解析成功,解析结果为[:：]?\s*(?P<ppt_struct>\{.*\})"
)
PAGE_RESULT_START_PATTERN = re.compile(
    r"\[(?P<task_id>[A-Za-z0-9]+)\]\[getPptStruct\]\s*获取文件解析结果,\s*page:\s*(?P<json_text>\{.*)"
)
CONTENT_PAGE_NUMBER_PATTERN = re.compile(r'"pageNumber"\s*:\s*(\d+)')


def parse_args() -> argparse.Namespace:
    """解析命令行参数。"""
    parser = argparse.ArgumentParser(description="根据 Excel 中的任务编号，从日志回填 pptStruct 与相关提示词输入。")
    parser.add_argument("--sheet-name", help="可选，指定工作表名称；默认使用 active sheet")
    parser.add_argument(
        "--mode",
        choices=("error-page", "full-file"),
        default="error-page",
        help="error-page 提取当前页；full-file 提取整份文件全部页的 pptStruct",
    )
    parser.add_argument(
        "--skip-non-empty",
        action="store_true",
        help="若 `日志中的pptstruct` 已有内容则跳过，避免覆盖已写入结果",
    )
    return parser.parse_args()


def normalize_header(value) -> str:
    """标准化表头文本，便于做精确列匹配。

    Args:
        value: 原始表头值。

    Returns:
        str: 去首尾空格并转小写后的表头。
    """
    if value is None:
        return ""
    return str(value).strip().lower()


def find_required_columns(sheet) -> Dict[str, int]:
    """按当前 Excel 结构定位所有必需列。

    Args:
        sheet: 当前工作表对象。

    Returns:
        Dict[str, int]: 逻辑列名到 Excel 列号的映射。
    """
    header_map: Dict[str, int] = {}
    for column_index in range(1, sheet.max_column + 1):
        normalized = normalize_header(sheet.cell(1, column_index).value)
        if normalized:
            header_map[normalized] = column_index

    required_aliases = {
        "task_id": (TASK_ID_HEADER,),
        "page_number": (PAGE_NUMBER_HEADER,),
        "ppt_struct": (PPT_STRUCT_HEADER,),
        "paddle_ocr": (PADDLE_OCR_HEADER,),
        "image_parse": (IMAGE_PARSE_HEADER,),
        "chart_parse": (CHART_PARSE_HEADER,),
    }
    result: Dict[str, int] = {}
    for key, aliases in required_aliases.items():
        column_index = 0
        for alias in aliases:
            column_index = header_map.get(normalize_header(alias), 0)
            if column_index:
                break
        if not column_index:
            raise AssertionError("Excel 缺少表头: {}".format(" / ".join(aliases)))
        result[key] = column_index
    return result


def list_log_files(log_source: Path) -> List[Path]:
    """列出待扫描日志文件。

    Args:
        log_source: 日志目录或单个日志文件路径。

    Returns:
        List[Path]: 实际参与扫描的日志文件列表。
    """
    resolved = log_source.expanduser().resolve()
    if resolved.is_file():
        return [resolved]
    if resolved.is_dir():
        return sorted(path for path in resolved.iterdir() if path.is_file() and path.name.startswith("my-pod") and path.suffix == ".log")
    raise FileNotFoundError("日志路径不存在: {}".format(resolved))


def normalize_page_number(value) -> Optional[str]:
    """将页码统一转为字符串。

    Args:
        value: 原始页码值。

    Returns:
        Optional[str]: 标准化后的页码字符串，缺失时返回 `None`。
    """
    if value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def overwrite_cell(sheet, row_index: int, column_index: int, value: str) -> None:
    """覆盖写入单元格内容。

    Args:
        sheet: 工作表对象。
        row_index: 行号。
        column_index: 列号。
        value: 待写入值。
    """
    cell = sheet.cell(row_index, column_index)
    if cell.value not in (None, ""):
        cell.value = None
    cell.value = value


def parse_json_block(lines: List[str], start_index: int, initial_text: str) -> Tuple[Optional[dict], int]:
    """从多行日志中增量解析一个完整 JSON 对象。

    Args:
        lines: 全部日志行。
        start_index: 当前 JSON 起始行。
        initial_text: 当前起始行中已截取出的 JSON 片段。

    Returns:
        Tuple[Optional[dict], int]:
            - 解析出的 JSON 对象，失败时为 `None`
            - JSON 结束时所在的日志行号
    """
    decoder = json.JSONDecoder()
    buffer = initial_text.rstrip("\n")
    current_index = start_index
    while True:
        try:
            parsed, end_index = decoder.raw_decode(buffer)
            if buffer[end_index:].strip():
                return None, current_index
            return parsed, current_index
        except json.JSONDecodeError:
            current_index += 1
            if current_index >= len(lines):
                return None, len(lines) - 1
            buffer += "\n" + lines[current_index].rstrip("\n")


def resolve_page_number_from_page_result(page_result: dict) -> Optional[str]:
    """从页解析结果中推断真实页码。

    Args:
        page_result: `获取文件解析结果` 对应的页级 JSON。

    Returns:
        Optional[str]: 推断出的页码字符串。
    """
    candidates = [
        page_result.get("page_number"),
        page_result.get("pageNumber"),
        page_result.get("page"),
    ]
    for candidate in candidates:
        normalized = normalize_page_number(candidate)
        if normalized:
            return normalized

    content_blocks_text = str(page_result.get("content_blocks") or "")
    match = CONTENT_PAGE_NUMBER_PATTERN.search(content_blocks_text)
    if match:
        return match.group(1)
    return None


def format_prompt_sections(page_result: dict) -> Dict[str, str]:
    """将页解析结果映射为后续审核使用的三段日志输入。

    Args:
        page_result: `获取文件解析结果` 对应的页级 JSON。

    Returns:
        Dict[str, str]: 三个日志字段的文本值。
    """
    return {
        "paddle_ocr": str(page_result.get("content_blocks") or ""),
        "image_parse": str(page_result.get("image_results") or ""),
        "chart_parse": str(page_result.get("chart_results") or ""),
    }


def build_log_indexes(log_paths: List[Path]) -> Dict[str, Dict[str, Dict[str, object]]]:
    """构建按 `taskid + 页码` 检索的日志索引。

    索引会同时保留：
    1. 单页 `pptStruct` 解析结果
    2. 该页提示词输入对应的 OCR / 图片 / 图表三段内容

    Args:
        log_paths: 待扫描日志文件列表。

    Returns:
        Dict[str, Dict[str, Dict[str, object]]]: 二级索引。
    """
    index: Dict[str, Dict[str, Dict[str, object]]] = {}
    for log_path in log_paths:
        lines = log_path.read_text(encoding="utf-8", errors="ignore").splitlines()
        line_index = 0
        while line_index < len(lines):
            line = lines[line_index]

            if "[getPptStruct]" in line and "获取文件解析结果, page:" in line:
                match = PAGE_RESULT_START_PATTERN.search(line)
                if match:
                    page_result, end_index = parse_json_block(lines, line_index, match.group("json_text"))
                    line_index = end_index
                    if isinstance(page_result, dict):
                        task_id = match.group("task_id")
                        page_number = resolve_page_number_from_page_result(page_result)
                        if page_number:
                            entry = index.setdefault(task_id, {}).setdefault(page_number, {})
                            entry.update(format_prompt_sections(page_result))
                line_index += 1
                continue

            if "[getPptStruct]" in line and "解析结果为" in line:
                match = PAGE_PPT_STRUCT_PATTERN.search(line)
                if match:
                    try:
                        parsed = json.loads(match.group("ppt_struct"))
                    except json.JSONDecodeError:
                        line_index += 1
                        continue
                    task_id = match.group("task_id")
                    page_number = match.group("page_number")
                    entry = index.setdefault(task_id, {}).setdefault(page_number, {})
                    entry["ppt_struct"] = parsed

            line_index += 1
    return index


def build_full_file_ppt_struct(page_map: Dict[str, Dict[str, object]]) -> str:
    """按页码顺序聚合整份文件的 `pptStruct`。

    Args:
        page_map: 单任务下的页级日志索引。

    Returns:
        str: 以页码为 key 的整份 `pptStruct` JSON 字符串。
    """
    ordered_pages = sorted(page_map.items(), key=lambda item: int(item[0]) if item[0].isdigit() else item[0])
    ordered_dict = {
        page_number: page_entry["ppt_struct"]
        for page_number, page_entry in ordered_pages
        if isinstance(page_entry, dict) and "ppt_struct" in page_entry
    }
    return json.dumps(ordered_dict, ensure_ascii=False)


def write_page_fields(sheet, row_index: int, column_map: Dict[str, int], page_entry: Dict[str, object]) -> None:
    """写入当前页的四个日志字段。

    Args:
        sheet: 工作表对象。
        row_index: 当前行号。
        column_map: 逻辑列名到 Excel 列号的映射。
        page_entry: 单页日志索引结果。
    """
    ppt_struct_value = page_entry.get("ppt_struct")
    if ppt_struct_value is None:
        overwrite_cell(sheet, row_index, column_map["ppt_struct"], "日志中未定位到该任务编号对应错误页码的pptStruct")
    else:
        overwrite_cell(sheet, row_index, column_map["ppt_struct"], json.dumps(ppt_struct_value, ensure_ascii=False))
    overwrite_cell(sheet, row_index, column_map["paddle_ocr"], str(page_entry.get("paddle_ocr") or ""))
    overwrite_cell(sheet, row_index, column_map["image_parse"], str(page_entry.get("image_parse") or ""))
    overwrite_cell(sheet, row_index, column_map["chart_parse"], str(page_entry.get("chart_parse") or ""))


def clear_aux_fields(sheet, row_index: int, column_map: Dict[str, int]) -> None:
    """清空三个辅助日志字段。

    Args:
        sheet: 工作表对象。
        row_index: 当前行号。
        column_map: 逻辑列名到 Excel 列号的映射。
    """
    overwrite_cell(sheet, row_index, column_map["paddle_ocr"], "")
    overwrite_cell(sheet, row_index, column_map["image_parse"], "")
    overwrite_cell(sheet, row_index, column_map["chart_parse"], "")


def main() -> int:
    """脚本主入口。

    Returns:
        int: 进程退出码。
    """
    args = parse_args()
    excel_path = DEFAULT_EXCEL_PATH.expanduser().resolve()
    log_paths = list_log_files(DEFAULT_LOG_SOURCE)

    workbook = load_workbook(excel_path)
    sheet = workbook[args.sheet_name] if args.sheet_name else workbook.active
    column_map = find_required_columns(sheet)
    log_index = build_log_indexes(log_paths)
    updated_rows = 0

    for row_index in range(2, sheet.max_row + 1):
        task_id = sheet.cell(row_index, column_map["task_id"]).value
        page_number = sheet.cell(row_index, column_map["page_number"]).value
        existing_value = sheet.cell(row_index, column_map["ppt_struct"]).value

        if args.skip_non_empty and existing_value:
            continue
        if not task_id:
            continue

        normalized_task_id = str(task_id).strip()
        page_map = log_index.get(normalized_task_id)
        if not isinstance(page_map, dict):
            overwrite_cell(sheet, row_index, column_map["ppt_struct"], "日志中未定位到该任务编号对应的pptStruct")
            clear_aux_fields(sheet, row_index, column_map)
            updated_rows += 1
            continue

        normalized_page_number = normalize_page_number(page_number)

        if args.mode == "full-file":
            overwrite_cell(sheet, row_index, column_map["ppt_struct"], build_full_file_ppt_struct(page_map))
            if normalized_page_number:
                page_entry = page_map.get(normalized_page_number)
                if isinstance(page_entry, dict):
                    overwrite_cell(sheet, row_index, column_map["paddle_ocr"], str(page_entry.get("paddle_ocr") or ""))
                    overwrite_cell(sheet, row_index, column_map["image_parse"], str(page_entry.get("image_parse") or ""))
                    overwrite_cell(sheet, row_index, column_map["chart_parse"], str(page_entry.get("chart_parse") or ""))
                else:
                    clear_aux_fields(sheet, row_index, column_map)
            else:
                clear_aux_fields(sheet, row_index, column_map)
            updated_rows += 1
            continue

        if not normalized_page_number:
            continue

        page_entry = page_map.get(normalized_page_number)
        if not isinstance(page_entry, dict):
            overwrite_cell(sheet, row_index, column_map["ppt_struct"], "日志中未定位到该任务编号对应错误页码的pptStruct")
            clear_aux_fields(sheet, row_index, column_map)
            updated_rows += 1
            continue

        write_page_fields(sheet, row_index, column_map, page_entry)
        updated_rows += 1

    workbook.save(excel_path)
    print("Excel 已更新: {}".format(excel_path))
    print("写入记录数: {}".format(updated_rows))
    return 0


if __name__ == "__main__":
    sys.exit(main())
