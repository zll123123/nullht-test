#!/usr/bin/env python3
"""对 Excel 中已存在的 `ppt原文` 与日志提取结果做审核与归因。

脚本职责：
1. 不负责 PPT 原文抽取。
2. 不负责从日志补提 `pptStruct`、`paddle OCR`、图片解析、图表解析。
3. 只消费 Excel 中已有列，完成两类输出：
   - `ppt原文` vs `日志中的pptstruct` 的语义审核结果
   - 基于审核结果与日志输入的错误原因归因，支持多标签，主因排第一

主要输入列：
1. `文件名称`
2. `页码`
3. `ppt原文`
4. `日志中的pptstruct`
5. `日志paddleocr识别结果`
6. `日志中的图片解析内容`
7. `日志中的图表解析内容`

主要输出列：
1. `LLM错误类型`
2. `LLM严重程度`
3. `LLM判断依据`
4. `LLM审核结果`
5. `原文问题`
6. `pptstruct问题`
7. `问题原因标签`（多个原因用 `/` 连接，主因排第一）
8. `问题原因说明`

执行模式：
1. `llm`
   - 基于 Excel 中现有内容做 LLM 审核与原因归因。
2. `cause-only`
   - 不重新执行 LLM 审核。
   - 直接复用表格里现有的 `LLM审核结果`、`pptstruct问题`，只重算
     `问题原因标签` 和 `问题原因说明`。
"""

import argparse
import json
import re
import sys
import time
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import requests
from loguru import logger
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


DEFAULT_EXCEL_PATH = Path(__file__).resolve().with_name("ppt原文抽取结果_获取pptstruct对比.xlsx")
DEFAULT_LOG_PATH = Path(__file__).resolve().with_name("review_pptstruct_semantics_to_excel.log")
DEFAULT_LLM_BASE_URL = "https://dashscope.aliyuncs.com/compatible-mode/v1"
DEFAULT_LLM_MODEL = "deepseek-v4-flash"
DEFAULT_LLM_API_KEY = "sk-433e97f12d014e2d9d74ce0669aa7155"
DEFAULT_LLM_TIMEOUT_SECONDS = 120
DEFAULT_LLM_MAX_RETRIES = 3

FILE_NAME_HEADER = "文件名称"
PAGE_NUMBER_HEADER = "页码"
PPT_STRUCT_HEADER = "日志中的pptstruct"
PPT_SOURCE_HEADER = "ppt原文"
LLM_ERROR_TYPE_HEADER = "LLM错误类型"
LLM_SEVERITY_HEADER = "LLM严重程度"
LLM_BASIS_HEADER = "LLM判断依据"
LLM_RESULT_HEADER = "LLM审核结果"
SOURCE_ISSUE_HEADER = "原文问题"
STRUCT_ISSUE_HEADER = "pptstruct问题"
ROOT_CAUSE_TAG_HEADER = "问题原因标签"
ROOT_CAUSE_REASON_HEADER = "问题原因说明"
PADDLE_OCR_HEADERS = ("日志paddleocr识别结果",)
IMAGE_PARSE_HEADER = "日志中的图片解析内容"
CHART_PARSE_HEADER = "日志中的图表解析内容"

SENTENCE_SEPARATOR_PATTERN = re.compile(r"[。！？!?；;]+|\n+")
WHITESPACE_PATTERN = re.compile(r"\s+")
ALIGNMENT_THRESHOLD = 0.55
MAX_REPORT_ITEMS = 5
MAX_TEXT_SNIPPET_LENGTH = 80
PASS_KEYWORDS = ("审核通过", "未发现明显问题", "未发现问题", "无问题")
OCR_ERROR_KEYWORDS = ("ocr", "识别", "文本块", "版面", "表格", "漏字", "错字", "文字", "字段缺失", "字段错误")
TOKEN_CHAR_CLASS = r"A-Za-z0-9\u4e00-\u9fff·\-ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩⅪⅫα-ωΑ-Ω"
TERM_PATTERN = re.compile(r"[\"'“”‘’]?([" + TOKEN_CHAR_CLASS + r"]{2,})[\"'“”‘’]?")
ISSUE_SPLIT_PATTERN = re.compile(r"[；;。]\s*|\s*(?:并且|且|同时|并)\s*")
REPLACE_PATTERNS = (
    re.compile(
        r"[\"'“”‘’]?(?P<wrong>[" + TOKEN_CHAR_CLASS + r"]{2,})[\"'“”‘’]?"
        r"与原文"
        r"[\"'“”‘’]?(?P<correct>[" + TOKEN_CHAR_CLASS + r"]{2,})[\"'“”‘’]?"
        r"不一致"
    ),
    re.compile(
        r"(?:将|把)?[\"'“”‘’]?(?P<correct>[" + TOKEN_CHAR_CLASS + r"]{2,})[\"'“”‘’]?"
        r"(?:错误)?(?:提取|识别|写|记|抽取)?(?:为|成)"
        r"[\"'“”‘’]?(?P<wrong>[" + TOKEN_CHAR_CLASS + r"]{2,})[\"'“”‘’]?"
    ),
)
MISSING_PATTERNS = (
    re.compile(r"(?:遗漏(?:了)?|缺少(?:了)?|缺失(?:了)?|未提取|未识别|未保留)(?P<term>[" + TOKEN_CHAR_CLASS + r"]{2,})"),
    re.compile(r"[\"'“”‘’]?(?P<term>[" + TOKEN_CHAR_CLASS + r"]{2,})[\"'“”‘’]?(?:未提取|未识别|被遗漏|被漏掉)"),
)
ROOT_CAUSE_PRIORITY = ("paddle VL识别错误", "图表或者图片错误", "pptstruct错误")
GENERIC_ISSUE_TERMS = {
    "pptstruct",
    "原文",
    "内容",
    "文本",
    "字段",
    "信息",
    "问题",
    "错误",
    "遗漏",
    "冗余",
    "重复",
    "表述",
    "偏差",
    "事实",
    "语义",
    "图片",
    "图表",
    "名称",
    "药物名称",
    "名词",
    "描述",
    "摘要",
    "总结",
    "正文",
    "与原文",
    "原文错误形式",
}


def parse_args() -> argparse.Namespace:
    """解析命令行参数。"""
    parser = argparse.ArgumentParser(description="基于 Excel 中已有的 PPT 原文与 pptStruct 文本做独立语义对比。")
    parser.add_argument("--sheet-name", help="可选，指定工作表名称；默认使用 active sheet")
    parser.add_argument(
        "--mode",
        choices=("llm", "cause-only"),
        default="llm",
        help="llm: 跑 LLM 审核与原因归因；cause-only: 仅重算归因列",
    )
    parser.add_argument(
        "--save-batch-size",
        type=int,
        default=10,
        help="每处理多少行保存一次 Excel，默认 10",
    )
    parser.add_argument(
        "--max-rows",
        type=int,
        help="最多处理多少条有效数据；不传则处理到文件末尾",
    )
    return parser.parse_args()


def setup_logging() -> None:
    """初始化控制台与文件日志。"""
    logger.remove()
    logger.add(sys.stdout, level="INFO", format="{time:YYYY-MM-DD HH:mm:ss} | {message}")
    logger.add(str(DEFAULT_LOG_PATH), level="INFO", encoding="utf-8", mode="w")


def normalize_name(name: str) -> str:
    """标准化文件名，便于后续扩展时做比对。"""
    normalized = name.strip().lower()
    normalized = normalized.replace(".pptx.pptx", ".pptx")
    normalized = normalized.replace(".ppt.ppt", ".ppt")
    normalized = WHITESPACE_PATTERN.sub("", normalized)
    return normalized


def normalize_text(text: str) -> str:
    """统一换行与空白格式。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [WHITESPACE_PATTERN.sub(" ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def compact_text(text: str) -> str:
    """压缩文本用于弱格式比对。"""
    return re.sub(r"[\s，,。！？!?；;:：、“”\"'‘’（）()\[\]【】·\-_/]", "", text).lower()


def shorten_text(text: str, max_length: int = MAX_TEXT_SNIPPET_LENGTH) -> str:
    """截断长文本片段，避免 diff 结果过长。"""
    text = normalize_text(text).replace("\n", " ")
    if len(text) <= max_length:
        return text
    return text[: max_length - 3] + "..."


def normalize_page_number(value) -> Optional[int]:
    """把页码统一转成整数。"""
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def collect_struct_texts(value) -> List[str]:
    """递归提取 pptStruct 中参与正文对比的 text 字段。"""
    texts: List[str] = []
    if isinstance(value, dict):
        text = value.get("text")
        if isinstance(text, str):
            normalized = normalize_text(text)
            if normalized:
                texts.append(normalized)
        for child_key, child_value in value.items():
            if child_key in {"text", "ppt_summary"}:
                continue
            texts.extend(collect_struct_texts(child_value))
    elif isinstance(value, list):
        for child in value:
            texts.extend(collect_struct_texts(child))
    return texts


def build_ppt_struct_text(ppt_struct_obj: Dict[str, object]) -> str:
    """将 pptStruct 中可比对文本按遍历顺序拼接。"""
    ordered_texts: List[str] = []
    seen = set()
    for text in collect_struct_texts(ppt_struct_obj):
        key = compact_text(text)
        if not key or key in seen:
            continue
        seen.add(key)
        ordered_texts.append(text)
    return normalize_text("\n".join(ordered_texts))


def split_sentences(text: str) -> List[str]:
    """按句号、分号和换行粗粒度切句。"""
    normalized = normalize_text(text)
    if not normalized:
        return []
    sentences: List[str] = []
    for chunk in SENTENCE_SEPARATOR_PATTERN.split(normalized):
        sentence = normalize_text(chunk)
        if not sentence:
            continue
        if len(compact_text(sentence)) <= 1:
            continue
        sentences.append(sentence)
    return sentences


def similarity(left: str, right: str) -> float:
    """计算两段文本的相似度。"""
    return SequenceMatcher(None, compact_text(left), compact_text(right)).ratio()


def build_best_match_map(sentences: List[str], candidates: List[str]) -> Dict[int, Tuple[Optional[int], float]]:
    """为每个句子寻找最接近的候选句。"""
    result: Dict[int, Tuple[Optional[int], float]] = {}
    for sentence_index, sentence in enumerate(sentences):
        best_index: Optional[int] = None
        best_score = 0.0
        for candidate_index, candidate in enumerate(candidates):
            score = similarity(sentence, candidate)
            if score > best_score:
                best_index = candidate_index
                best_score = score
        result[sentence_index] = (best_index, best_score)
    return result


def summarize_char_diff(source_text: str, target_text: str, max_items: int = 3) -> str:
    """对已对齐句子生成字符级差异摘要。"""
    matcher = SequenceMatcher(None, source_text, target_text)
    parts: List[str] = []
    for tag, source_start, source_end, target_start, target_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        source_fragment = source_text[source_start:source_end]
        target_fragment = target_text[target_start:target_end]
        if tag == "replace":
            parts.append("替换[{} -> {}]".format(shorten_text(source_fragment, 20), shorten_text(target_fragment, 20)))
        elif tag == "delete":
            parts.append("缺失[{}]".format(shorten_text(source_fragment, 20)))
        elif tag == "insert":
            parts.append("新增[{}]".format(shorten_text(target_fragment, 20)))
        if len(parts) >= max_items:
            break
    return "；".join(parts) if parts else "存在字符级差异"


def short_join(items: List[str], empty_text: str) -> str:
    """限制摘要项数量并拼接结果。"""
    if not items:
        return empty_text
    return "；".join(items[:MAX_REPORT_ITEMS])


def find_header_columns(sheet) -> Dict[str, int]:
    """读取首行表头到列号的映射。"""
    header_map: Dict[str, int] = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str):
            header = cell.value.strip()
            if header:
                header_map[header] = cell.column
    return header_map


def find_first_existing_column(header_map: Dict[str, int], headers: Tuple[str, ...]) -> Optional[int]:
    """从一组候选表头中找到首个存在的列。"""
    for header in headers:
        column_index = header_map.get(header)
        if column_index is not None:
            return column_index
    return None


def ensure_output_column(sheet, header_map: Dict[str, int], header_name: str, preferred_index: int) -> int:
    """确保输出列存在，不存在时追加创建。"""
    column_index = header_map.get(header_name)
    if column_index is None:
        column_index = preferred_index
        sheet.cell(1, column_index).value = header_name
        header_map[header_name] = column_index
    return column_index


def require_columns(header_map: Dict[str, int], headers: List[str]) -> Dict[str, int]:
    """校验必需输入列全部存在。"""
    missing = [header for header in headers if header not in header_map]
    if missing:
        raise RuntimeError("Excel 缺少表头: {}".format(", ".join(missing)))
    return {header: header_map[header] for header in headers}


def require_column_aliases(header_map: Dict[str, int], column_aliases: Dict[str, Tuple[str, ...]]) -> Dict[str, int]:
    """校验一组带别名的输入列全部存在。"""
    result: Dict[str, int] = {}
    missing: List[str] = []
    for key, aliases in column_aliases.items():
        column_index = find_first_existing_column(header_map, aliases)
        if column_index is None:
            missing.append(" / ".join(aliases))
            continue
        result[key] = column_index
    if missing:
        raise RuntimeError("Excel 缺少表头: {}".format(", ".join(missing)))
    return result


def build_llm_prompt(source_text: str, struct_text: str) -> str:
    """构造主审核提示词。"""
    return """你是一名严格的医学内容审核助手。
请基于给定的PPT原文和pptStruct文本，重点检查以下医学关键信息是否在pptStruct中正确呈现，不得有任何错误：
1. 临床与研究核心数据：研究对象(PICO)、样本量、分组/盲法；注册号(NCT)、方案版本/日期、伦理批件号；年龄/性别/种族、疾病分期(如TNM、Child-Pugh、Rai等)、基线评分(如NIHSS、TNM)、关键体征/合并症；主要/次要终点(如OS、PFS)；置信区间、P值、HR/RR等统计结果；所有不良事件(AE/SAE)及其归因、严重程度、处理措施；专有名词(药品/器械名)拼写、缩写；数值单位(如mmol/L误为mol/L)。
2. 背景与合规信息：仅检查pptStruct是否正确提取了原文中的以下内容——作者/机构、参考文献完整出处、利益冲突声明、超说明书用药等。
3. 逻辑与结论：核心结论(Take-home message)是否准确；叙述逻辑是否一致；关键支撑数据是否被误删或扭曲。

特别说明：
- `image_elements` 下的 `text` 字段是从原文中提取出的文本，需要参与审核，并判断其内容是否正确、是否与原文一致。
- 仅关注内容层面的语义、事实、表述等问题。

判断问题类型、严重程度、依据，并分别指出原文问题和pptStruct问题。

输出要求：
1. 仅输出JSON，不要输出Markdown代码块。
2. JSON字段固定为：error_type, severity, basis, result, source_issue, struct_issue。
3. error_type 只针对“pptStruct问题”进行判断，不评价原文问题；从以下枚举中选一个或多个并用顿号连接：语义遗漏、语义冗余、事实错误、表述偏差、无、无法判断。
4. severity 只针对“pptStruct问题”的严重程度进行判断，只能是：高、中、低。（医学关键信息错误/遗漏 -> 高；非关键但影响理解 -> 中；轻微不影响医学判断 -> 低）
5. 如果审核通过或未发现明显问题，basis 必须返回空字符串。
6. 只有判断存在不一致或风险时，basis 才填写中文简洁说明判断依据（例如：遗漏了主要终点OS值、单位错误）。
7. result 用中文总结审核意见。
8. source_issue 只填写“原文存在的问题”，没有则填“无”。
9. struct_issue 只填写“pptStruct存在的问题”，没有则填“无”。
10. 如果 struct_issue = 无，那么 error_type 必须是“无”，severity 必须是“低”。
11. 如果只有原文问题、而 pptStruct 没有问题，那么 error_type 仍然填“无”，severity 仍然填“低”。
12. “提取重复”只在“同一个字段内同一内容重复出现多次”时成立；如果同一内容出现在不同字段中（例如正文字段和声明字段各出现一次），不算提取重复。
13. ppt_summary 是对整页 PPT 的总结，可以简单检查 summary 是否合理，但 summary 本身不参与正文对比，也不要因为 summary 和正文字段重复就判定提取重复。

PPT原文：
{source_text}

pptStruct文本：
{struct_text}
""".format(source_text=source_text, struct_text=struct_text)


def build_root_cause_prompt(
    llm_result_text: str,
    struct_issue_text: str,
    paddle_ocr_text: str,
    image_parse_text: str,
    chart_parse_text: str,
) -> str:
    """构造多标签错误原因归因提示词。"""
    return """你是一名严格的问题归因助手。
请先从 `pptstruct审核结果` 中提取每一条具体的 `pptstruct问题`，然后对每条问题单独归因，再汇总整体标签。

可选分类只允许以下 4 个值：
1. paddle VL识别错误
2. 图表或者图片错误
3. pptstruct错误
4. 无

判断规则：
1. 只有当 `LLM审核结果` 明确表示无问题、审核通过、未发现问题，且 `pptstruct问题` 为 `无` 时，categories 才返回 `["无"]`，reason 返回空字符串。
2. 先把 `pptstruct问题` 拆成若干具体问题，每条问题归入以下三类之一：
   - `遗漏`：pptStruct 少了原文应有的内容
   - `多出内容`：pptStruct 比原文多了不该有的内容
   - `内容错误`：pptStruct 把内容提错了、写错了、值错了、实体错了
3. 对于 `遗漏`：
   - 先检查被遗漏内容是否出现在 `日志paddleocr识别结果`、`日志中的图片解析内容`、`日志中的图表解析内容` 中。
   - 若任一上游结果中已经出现该内容，而最终 pptStruct 未保留，则归因为 `pptstruct错误`。
   - 若上游都未出现，再判断该内容理论上应来自哪里：
     - 正文文字、表格文字、参考文献文字、页内文本 -> `paddle VL识别错误`
     - 图片语义、图表语义、图片/图表说明 -> `图表或者图片错误`
4. 对于 `多出内容`：
   - 检查多出的内容最早出现在哪个上游结果中。
   - 若 `日志paddleocr识别结果` 已出现该多余内容，则归因为 `paddle VL识别错误`。
   - 若 `日志中的图片解析内容` 或 `日志中的图表解析内容` 已出现该多余内容，则归因为 `图表或者图片错误`。
   - 若上游都未出现，而只有 pptStruct 出现，则归因为 `pptstruct错误`。
5. 对于 `内容错误`：
   - 只追踪错误版本的内容，不追踪正确版本内容。
   - 若错误内容最早出现在 `日志paddleocr识别结果`，归因为 `paddle VL识别错误`。
   - 若错误内容最早出现在 `日志中的图片解析内容` 或 `日志中的图表解析内容`，归因为 `图表或者图片错误`。
   - 若上游都未出现错误内容，而 pptStruct 出现错误内容，则归因为 `pptstruct错误`。
6. `日志中的图片解析内容`、`日志中的图表解析内容` 里大量描述性语句和多模态噪声不能直接作为证据；只有出现可直接核对的明确文本证据时，才能作为归因依据。
7. 允许多选，但整体 `categories` 必须按固定优先级排序：`paddle VL识别错误` > `图表或者图片错误` > `pptstruct错误`。
8. `reason` 必须逐条写清楚，每条都使用以下格式：
   - `xxx归因为xxxx，原因是xxxx`
   - 如果某条问题无法明确归因，写成：`xxx无法归因，原因是xxxx`
   - 多条问题之间用 `；` 连接。
9. 仅输出 JSON，不要输出 Markdown。

JSON 字段固定为：
- categories
- reason

LLM审核结果：
{llm_result_text}

pptstruct问题：
{struct_issue_text}

日志paddleocr识别结果：
{paddle_ocr_text}

日志中的图片解析内容：
{image_parse_text}

日志中的图表解析内容：
{chart_parse_text}
""".format(
        llm_result_text=llm_result_text,
        struct_issue_text=struct_issue_text,
        paddle_ocr_text=paddle_ocr_text,
        image_parse_text=image_parse_text,
        chart_parse_text=chart_parse_text,
    )


def normalize_llm_json(content: str) -> Dict[str, str]:
    """解析并兜底主审核模型输出。"""
    normalized = content.strip()
    normalized = re.sub(r"<think>.*?</think>", "", normalized, flags=re.DOTALL).strip()
    if normalized.startswith("```"):
        normalized = normalized.strip("`")
        if normalized.startswith("json"):
            normalized = normalized[4:].strip()
    parsed = json.loads(normalized)
    source_issue = str(parsed.get("source_issue") or "无")
    struct_issue = str(parsed.get("struct_issue") or "无")
    error_type = str(parsed.get("error_type") or "无法判断")
    severity = str(parsed.get("severity") or "无法判断")
    if struct_issue == "无":
        error_type = "无"
        severity = "低"
    return {
        "error_type": error_type,
        "severity": severity,
        "basis": str(parsed.get("basis") or ""),
        "result": str(parsed.get("result") or ""),
        "source_issue": source_issue,
        "struct_issue": struct_issue,
    }


def normalize_root_cause_json(content: str) -> Dict[str, str]:
    """解析并兜底多标签原因归因模型输出。"""
    normalized = content.strip()
    normalized = re.sub(r"<think>.*?</think>", "", normalized, flags=re.DOTALL).strip()
    if normalized.startswith("```"):
        normalized = normalized.strip("`")
        if normalized.startswith("json"):
            normalized = normalized[4:].strip()
    parsed = json.loads(normalized)
    categories_raw = parsed.get("categories")
    reason = str(parsed.get("reason") or "")
    allowed_categories = {"paddle VL识别错误", "图表或者图片错误", "pptstruct错误", "无"}
    categories: List[str] = []
    if isinstance(categories_raw, list):
        for item in categories_raw:
            text = str(item or "").strip()
            if text in allowed_categories and text not in categories:
                categories.append(text)
    elif isinstance(categories_raw, str):
        for item in re.split(r"[、/,，\s]+", categories_raw):
            text = item.strip()
            if text in allowed_categories and text not in categories:
                categories.append(text)
    if not categories:
        categories = ["无"]
        reason = ""
    if "无" in categories:
        categories = ["无"]
        reason = ""
    return {"categories": categories, "reason": reason}


def call_llm_review(source_text: str, struct_text: str) -> Dict[str, str]:
    """调用主审核模型，判断 pptStruct 语义问题。"""
    if not DEFAULT_LLM_API_KEY or "请在这里填写实际LLM_API_KEY" in DEFAULT_LLM_API_KEY:
        return {
            "error_type": "LLM未配置",
            "severity": "无法判断",
            "basis": "请先在代码中配置 DEFAULT_LLM_API_KEY",
            "result": "未执行LLM审核",
            "source_issue": "无",
            "struct_issue": "无",
        }
    payload = {
        "model": DEFAULT_LLM_MODEL,
        "temperature": 0.1,
        "messages": [
            {"role": "system", "content": "你只返回JSON。"},
            {"role": "user", "content": build_llm_prompt(source_text, struct_text)},
        ],
    }
    last_error: Optional[Exception] = None
    for attempt in range(1, DEFAULT_LLM_MAX_RETRIES + 1):
        try:
            response = requests.post(
                DEFAULT_LLM_BASE_URL.rstrip("/") + "/chat/completions",
                headers={
                    "Authorization": "Bearer {}".format(DEFAULT_LLM_API_KEY),
                    "Content-Type": "application/json",
                },
                json=payload,
                timeout=DEFAULT_LLM_TIMEOUT_SECONDS,
            )
            if response.status_code in {429, 500, 502, 503, 504, 529}:
                raise RuntimeError(
                    "LLM 服务暂时不可用(status={}): {}".format(response.status_code, shorten_text(response.text, 300))
                )
            response.raise_for_status()
            response_json = response.json()
            content = response_json["choices"][0]["message"]["content"]
            return normalize_llm_json(content)
        except Exception as exc:
            last_error = exc
            if attempt >= DEFAULT_LLM_MAX_RETRIES:
                break
            logger.warning("LLM 调用第 {} 次失败，准备重试: {}", attempt, exc)
            time.sleep(attempt)
    raise RuntimeError(str(last_error))


def call_llm_root_cause(
    llm_result_text: str,
    struct_issue_text: str,
    paddle_ocr_text: str,
    image_parse_text: str,
    chart_parse_text: str,
) -> Dict[str, str]:
    """调用原因归因模型，给出多标签错误来源分类。"""
    payload = {
        "model": DEFAULT_LLM_MODEL,
        "temperature": 0.1,
        "messages": [
            {"role": "system", "content": "你只返回JSON。"},
            {
                "role": "user",
                "content": build_root_cause_prompt(
                    llm_result_text=llm_result_text,
                    struct_issue_text=struct_issue_text,
                    paddle_ocr_text=paddle_ocr_text,
                    image_parse_text=image_parse_text,
                    chart_parse_text=chart_parse_text,
                ),
            },
        ],
    }
    last_error: Optional[Exception] = None
    for attempt in range(1, DEFAULT_LLM_MAX_RETRIES + 1):
        try:
            response = requests.post(
                DEFAULT_LLM_BASE_URL.rstrip("/") + "/chat/completions",
                headers={
                    "Authorization": "Bearer {}".format(DEFAULT_LLM_API_KEY),
                    "Content-Type": "application/json",
                },
                json=payload,
                timeout=DEFAULT_LLM_TIMEOUT_SECONDS,
            )
            if response.status_code in {429, 500, 502, 503, 504, 529}:
                raise RuntimeError(
                    "LLM 服务暂时不可用(status={}): {}".format(response.status_code, shorten_text(response.text, 300))
                )
            response.raise_for_status()
            response_json = response.json()
            content = response_json["choices"][0]["message"]["content"]
            return normalize_root_cause_json(content)
        except Exception as exc:
            last_error = exc
            if attempt >= DEFAULT_LLM_MAX_RETRIES:
                break
            logger.warning("问题原因分析第 {} 次失败，准备重试: {}", attempt, exc)
            time.sleep(attempt)
    raise RuntimeError(str(last_error))


def run_llm_compare(source_text: str, struct_text: str) -> Dict[str, str]:
    """运行主审核 LLM。"""
    return call_llm_review(source_text, struct_text)


def looks_like_pass_result(llm_result_text: str) -> bool:
    """判断当前审核结果是否明确表达为通过或无问题。"""
    normalized = normalize_text(llm_result_text).lower()
    if not normalized:
        return False
    return any(keyword in normalized for keyword in PASS_KEYWORDS)


def is_empty_llm_result(llm_result_text: str) -> bool:
    """判断 LLM 审核结果是否为空。"""
    return not normalize_text(llm_result_text)


def contains_any_keyword(text: str, keywords: Tuple[str, ...]) -> bool:
    """判断文本是否命中任一关键词。"""
    normalized = normalize_text(text).lower()
    return any(keyword.lower() in normalized for keyword in keywords)


def extract_candidate_terms(text: str) -> List[str]:
    """从问题描述中提取候选错词，用于判断是否为 OCR 继承错误。"""
    candidates: List[str] = []
    for match in TERM_PATTERN.findall(normalize_text(text)):
        normalized = match.strip()
        if len(normalized) < 2:
            continue
        if normalized.lower() in GENERIC_ISSUE_TERMS:
            continue
        if normalized not in candidates:
            candidates.append(normalized)
    return candidates


def split_issue_clauses(text: str) -> List[str]:
    """将问题描述切分为多个子问题。"""
    clauses: List[str] = []
    for clause in ISSUE_SPLIT_PATTERN.split(normalize_text(text)):
        normalized = clause.strip("，,：: ")
        if normalized:
            clauses.append(normalized)
    return clauses


def parse_struct_issue_items(struct_issue_text: str) -> List[Dict[str, object]]:
    """把 `pptstruct问题` 解析为结构化子问题。

    返回三类子问题：
    1. `replace`: 替换型错误，只追错误词最早出现位置
    2. `missing`: 遗漏型错误，检查目标词在哪一层被丢失
    3. `generic`: 无法结构化时的兜底问题
    """
    items: List[Dict[str, object]] = []
    for clause in split_issue_clauses(struct_issue_text):
        matched = False
        for pattern in REPLACE_PATTERNS:
            result = pattern.search(clause)
            if not result:
                continue
            correct_term = result.group("correct").strip()
            wrong_term = result.group("wrong").strip()
            if correct_term and wrong_term and correct_term != wrong_term:
                items.append(
                    {
                        "type": "replace",
                        "correct_term": correct_term,
                        "wrong_term": wrong_term,
                        "raw": clause,
                    }
                )
                matched = True
                break
        if matched:
            continue
        for pattern in MISSING_PATTERNS:
            result = pattern.search(clause)
            if not result:
                continue
            missing_term = result.group("term").strip()
            if missing_term and missing_term not in GENERIC_ISSUE_TERMS:
                items.append({"type": "missing", "missing_term": missing_term, "raw": clause})
                matched = True
                break
        if matched:
            continue
        generic_terms = extract_candidate_terms(clause)
        if generic_terms:
            items.append({"type": "generic", "terms": generic_terms, "raw": clause})

    if items:
        return items
    generic_terms = extract_candidate_terms(struct_issue_text)
    if generic_terms:
        return [{"type": "generic", "terms": generic_terms, "raw": normalize_text(struct_issue_text)}]
    return []


def match_terms_in_text(terms: List[str], target_text: str) -> List[str]:
    """返回在目标文本中命中的候选词。"""
    target_compact = compact_text(target_text)
    matched_terms: List[str] = []
    for term in terms:
        term_compact = compact_text(term)
        if not term_compact:
            continue
        if term_compact in target_compact and term not in matched_terms:
            matched_terms.append(term)
    return matched_terms


def append_root_cause_reason(categories: List[str], reasons: List[str], category: str, reason: str) -> None:
    """追加归因结果，保持标签去重。"""
    if category not in categories:
        categories.append(category)
    reasons.append(reason)


def build_reason_entry(issue_text: str, category: Optional[str], detail: str) -> str:
    """构造统一的归因说明文本。"""
    normalized_issue = normalize_text(issue_text) or "该问题"
    if category:
        return "{}归因为{}，原因是{}。".format(normalized_issue, category, detail)
    return "{}无法归因，原因是{}。".format(normalized_issue, detail)


def classify_replace_issue(
    issue_item: Dict[str, object],
    categories: List[str],
    reasons: List[str],
    paddle_ocr_text: str,
    image_chart_text: str,
    struct_text: str,
) -> None:
    """对替换型错误做归因，只追错误词。"""
    wrong_term = str(issue_item["wrong_term"])
    issue_text = str(issue_item.get("raw") or wrong_term)
    if match_terms_in_text([wrong_term], paddle_ocr_text):
        append_root_cause_reason(
            categories,
            reasons,
            "paddle VL识别错误",
            build_reason_entry(issue_text, "paddle VL识别错误", "替换错误“{}”最早出现在日志paddleocr识别结果中".format(wrong_term)),
        )
        return
    if match_terms_in_text([wrong_term], image_chart_text):
        append_root_cause_reason(
            categories,
            reasons,
            "图表或者图片错误",
            build_reason_entry(issue_text, "图表或者图片错误", "替换错误“{}”最早出现在图片/图表解析内容中".format(wrong_term)),
        )
        return
    if match_terms_in_text([wrong_term], struct_text):
        append_root_cause_reason(
            categories,
            reasons,
            "pptstruct错误",
            build_reason_entry(issue_text, "pptstruct错误", "替换错误“{}”未在上游日志中命中，属于pptstruct阶段引入".format(wrong_term)),
        )
        return
    append_root_cause_reason(
        categories,
        reasons,
        "pptstruct错误",
        build_reason_entry(issue_text, "pptstruct错误", "替换错误“{}”未找到明确上游来源，按pptstruct错误归因".format(wrong_term)),
    )


def classify_missing_issue(
    issue_item: Dict[str, object],
    categories: List[str],
    reasons: List[str],
    paddle_ocr_text: str,
    image_chart_text: str,
) -> None:
    """对遗漏型错误做归因。"""
    missing_term = str(issue_item["missing_term"])
    issue_text = str(issue_item.get("raw") or missing_term)
    if match_terms_in_text([missing_term], paddle_ocr_text):
        append_root_cause_reason(
            categories,
            reasons,
            "pptstruct错误",
            build_reason_entry(issue_text, "pptstruct错误", "遗漏项“{}”已在日志paddleocr识别结果中出现，但最终pptStruct未保留".format(missing_term)),
        )
        return
    if match_terms_in_text([missing_term], image_chart_text):
        append_root_cause_reason(
            categories,
            reasons,
            "pptstruct错误",
            build_reason_entry(issue_text, "pptstruct错误", "遗漏项“{}”已在图片/图表解析内容中出现，但最终pptStruct未保留".format(missing_term)),
        )
        return
    append_root_cause_reason(
        categories,
        reasons,
        "paddle VL识别错误",
        build_reason_entry(issue_text, "paddle VL识别错误", "遗漏项“{}”未在上游识别结果中出现，按上游识别缺失归因".format(missing_term)),
    )


def classify_generic_issue(
    issue_item: Dict[str, object],
    categories: List[str],
    reasons: List[str],
    paddle_ocr_text: str,
    image_chart_text: str,
    struct_text: str,
) -> None:
    """对未结构化的问题做兜底归因。"""
    terms = [term for term in issue_item.get("terms", []) if isinstance(term, str)]
    issue_text = str(issue_item.get("raw") or "该问题")
    if not terms:
        append_root_cause_reason(
            categories,
            reasons,
            "pptstruct错误",
            build_reason_entry(issue_text, "pptstruct错误", "问题描述过于泛化，按pptstruct错误归因"),
        )
        return
    struct_terms = match_terms_in_text(terms, struct_text)
    unresolved_terms = list(struct_terms or terms)
    paddle_terms = match_terms_in_text(unresolved_terms, paddle_ocr_text)
    if paddle_terms:
        append_root_cause_reason(
            categories,
            reasons,
            "paddle VL识别错误",
            build_reason_entry(issue_text, "paddle VL识别错误", "日志paddleocr识别结果已出现问题词：{}".format("、".join(paddle_terms[:5]))),
        )
        unresolved_terms = [term for term in unresolved_terms if term not in paddle_terms]
    image_chart_terms = match_terms_in_text(unresolved_terms, image_chart_text)
    if image_chart_terms:
        append_root_cause_reason(
            categories,
            reasons,
            "图表或者图片错误",
            build_reason_entry(issue_text, "图表或者图片错误", "图片/图表解析内容中出现问题词：{}".format("、".join(image_chart_terms[:5]))),
        )
        unresolved_terms = [term for term in unresolved_terms if term not in image_chart_terms]
    if unresolved_terms or not categories:
        append_root_cause_reason(
            categories,
            reasons,
            "pptstruct错误",
            build_reason_entry(issue_text, "pptstruct错误", "剩余未在上游日志中命中的问题词按pptstruct错误归因：{}".format("、".join(unresolved_terms[:5]))),
        )


def build_root_cause_by_terms(
    struct_issue_text: str,
    paddle_ocr_text: str,
    image_parse_text: str,
    chart_parse_text: str,
    struct_text: str,
) -> Dict[str, object]:
    """按结构化子问题构建归因结果。

    归因顺序固定为：
    1. `paddle VL识别错误`
    2. `图表或者图片错误`
    3. `pptstruct错误`
    """
    issue_items = parse_struct_issue_items(struct_issue_text)
    if not issue_items:
        return {
            "categories": ["pptstruct错误"],
            "reason": build_reason_entry(struct_issue_text, "pptstruct错误", "未从pptstruct问题中提取到结构化子问题，剩余问题按pptstruct错误归因"),
        }
    categories: List[str] = []
    reasons: List[str] = []
    image_chart_text = "\n".join(part for part in (image_parse_text, chart_parse_text) if normalize_text(part))
    for issue_item in issue_items:
        issue_type = str(issue_item.get("type") or "")
        if issue_type == "replace":
            classify_replace_issue(issue_item, categories, reasons, paddle_ocr_text, image_chart_text, struct_text)
            continue
        if issue_type == "missing":
            classify_missing_issue(issue_item, categories, reasons, paddle_ocr_text, image_chart_text)
            continue
        classify_generic_issue(issue_item, categories, reasons, paddle_ocr_text, image_chart_text, struct_text)
    ordered_categories = [category for category in ROOT_CAUSE_PRIORITY if category in categories]
    if not ordered_categories:
        ordered_categories = ["pptstruct错误"]
    return {"categories": ordered_categories, "reason": "；".join(reasons)}


def detect_root_cause_by_rules(
    llm_result_text: str,
    struct_issue_text: str,
    image_parse_text: str,
    chart_parse_text: str,
    paddle_ocr_text: str,
    struct_text: str,
) -> Optional[Dict[str, str]]:
    """执行最小化规则短路，复杂情况交给归因 LLM。"""
    if looks_like_pass_result(llm_result_text):
        return {"categories": ["无"], "reason": ""}

    if normalize_text(str(struct_issue_text or "")) == "无":
        return {"categories": ["无"], "reason": ""}

    return None


def format_root_cause_text(categories: List[str], reason: str) -> str:
    """将多标签归因结果格式化为标签列文本。"""
    if not categories or categories == ["无"]:
        return "无"
    if categories == ["未分析"]:
        return "未分析"
    return "/".join(categories)


def format_root_cause_reason(categories: List[str], reason: str) -> str:
    """将多标签归因结果格式化为说明列文本。"""
    if not categories or categories == ["无"]:
        return "无"
    if categories == ["未分析"]:
        return reason or "未分析"
    return reason or ""


def save_if_needed(
    workbook,
    excel_path: Path,
    processed_rows: int,
    batch_size: int,
    current_row_index: int,
    file_name: str,
    page_number: Optional[int],
) -> None:
    """按批次保存 Excel，并打印当前落盘进度。"""
    if processed_rows % batch_size == 0:
        logger.info(
            "准备保存 Excel: processed_rows={} current_row={} file={} page={}",
            processed_rows,
            current_row_index,
            file_name,
            page_number,
        )
        workbook.save(excel_path)
        logger.info(
            "批量保存完成: processed_rows={} current_row={} file={} page={} excel={}",
            processed_rows,
            current_row_index,
            file_name,
            page_number,
            excel_path,
        )


def main() -> int:
    """脚本主入口。"""
    setup_logging()
    args = parse_args()
    if args.save_batch_size <= 0:
        raise RuntimeError("--save-batch-size 必须大于 0")
    excel_path = DEFAULT_EXCEL_PATH.expanduser().resolve()
    workbook = load_workbook(excel_path)
    sheet = workbook[args.sheet_name] if args.sheet_name else workbook.active
    header_map = find_header_columns(sheet)
    if args.mode == "llm":
        required_input_columns = require_columns(
            header_map,
            [FILE_NAME_HEADER, PAGE_NUMBER_HEADER, PPT_STRUCT_HEADER, PPT_SOURCE_HEADER],
        )
    else:
        required_input_columns = require_columns(
            header_map,
            [FILE_NAME_HEADER, PAGE_NUMBER_HEADER, PPT_STRUCT_HEADER, LLM_RESULT_HEADER, STRUCT_ISSUE_HEADER],
        )
    root_cause_input_columns = require_column_aliases(
        header_map,
        {
            "paddle_ocr": PADDLE_OCR_HEADERS,
            "image_parse": (IMAGE_PARSE_HEADER,),
            "chart_parse": (CHART_PARSE_HEADER,),
        },
    )
    llm_error_type_column = ensure_output_column(sheet, header_map, LLM_ERROR_TYPE_HEADER, sheet.max_column + 1)
    llm_severity_column = ensure_output_column(sheet, header_map, LLM_SEVERITY_HEADER, max(sheet.max_column + 1, llm_error_type_column + 1))
    llm_basis_column = ensure_output_column(sheet, header_map, LLM_BASIS_HEADER, max(sheet.max_column + 1, llm_severity_column + 1))
    llm_result_column = ensure_output_column(sheet, header_map, LLM_RESULT_HEADER, max(sheet.max_column + 1, llm_basis_column + 1))
    source_issue_column = ensure_output_column(sheet, header_map, SOURCE_ISSUE_HEADER, max(sheet.max_column + 1, llm_result_column + 1))
    struct_issue_column = ensure_output_column(sheet, header_map, STRUCT_ISSUE_HEADER, max(sheet.max_column + 1, source_issue_column + 1))
    root_cause_tag_column = ensure_output_column(sheet, header_map, ROOT_CAUSE_TAG_HEADER, max(sheet.max_column + 1, struct_issue_column + 1))
    root_cause_reason_column = ensure_output_column(
        sheet,
        header_map,
        ROOT_CAUSE_REASON_HEADER,
        max(sheet.max_column + 1, root_cause_tag_column + 1),
    )
    sheet.column_dimensions[get_column_letter(llm_error_type_column)].width = 20
    sheet.column_dimensions[get_column_letter(llm_severity_column)].width = 12
    sheet.column_dimensions[get_column_letter(llm_basis_column)].width = 60
    sheet.column_dimensions[get_column_letter(llm_result_column)].width = 40
    sheet.column_dimensions[get_column_letter(source_issue_column)].width = 40
    sheet.column_dimensions[get_column_letter(struct_issue_column)].width = 40
    sheet.column_dimensions[get_column_letter(root_cause_tag_column)].width = 30
    sheet.column_dimensions[get_column_letter(root_cause_reason_column)].width = 60
    processed_rows = 0

    for row_index in range(2, sheet.max_row + 1):
        if args.max_rows is not None and processed_rows >= args.max_rows:
            logger.info("达到最大处理条数，提前停止: max_rows={}", args.max_rows)
            break
        file_name = sheet.cell(row_index, required_input_columns[FILE_NAME_HEADER]).value
        page_number = normalize_page_number(sheet.cell(row_index, required_input_columns[PAGE_NUMBER_HEADER]).value)
        ppt_struct_raw = sheet.cell(row_index, required_input_columns[PPT_STRUCT_HEADER]).value
        source_text_raw = (
            sheet.cell(row_index, required_input_columns[PPT_SOURCE_HEADER]).value
            if args.mode == "llm"
            else None
        )
        paddle_ocr_text = normalize_text(str(sheet.cell(row_index, root_cause_input_columns["paddle_ocr"]).value or ""))
        image_parse_text = normalize_text(str(sheet.cell(row_index, root_cause_input_columns["image_parse"]).value or ""))
        chart_parse_text = normalize_text(str(sheet.cell(row_index, root_cause_input_columns["chart_parse"]).value or ""))
        if not file_name or page_number is None or not ppt_struct_raw:
            continue
        logger.info("processing row={} file={} page={} mode={}", row_index, file_name, page_number, args.mode)

        try:
            ppt_struct_obj = json.loads(str(ppt_struct_raw))
        except Exception as exc:
            sheet.cell(row_index, llm_error_type_column).value = "无法判断"
            sheet.cell(row_index, llm_severity_column).value = "无法判断"
            sheet.cell(row_index, llm_basis_column).value = "pptStruct 解析失败: {}".format(exc)
            sheet.cell(row_index, llm_result_column).value = "未执行LLM审核"
            sheet.cell(row_index, source_issue_column).value = "无"
            sheet.cell(row_index, struct_issue_column).value = "pptStruct 解析失败: {}".format(exc)
            sheet.cell(row_index, root_cause_tag_column).value = "无"
            sheet.cell(row_index, root_cause_reason_column).value = "无"
            processed_rows += 1
            save_if_needed(workbook, excel_path, processed_rows, args.save_batch_size, row_index, str(file_name), page_number)
            continue

        if args.mode == "cause-only":
            llm_result = {
                "error_type": str(sheet.cell(row_index, llm_error_type_column).value or ""),
                "severity": str(sheet.cell(row_index, llm_severity_column).value or ""),
                "basis": str(sheet.cell(row_index, llm_basis_column).value or ""),
                "result": str(sheet.cell(row_index, header_map[LLM_RESULT_HEADER]).value or ""),
                "source_issue": str(sheet.cell(row_index, source_issue_column).value or ""),
                "struct_issue": str(sheet.cell(row_index, header_map[STRUCT_ISSUE_HEADER]).value or ""),
            }
        else:
            source_text = normalize_text(str(source_text_raw or ""))
            if not source_text:
                llm_result = {
                    "error_type": "无法判断",
                    "severity": "无法判断",
                    "basis": "PPT 原文为空，无法对比",
                    "result": "未执行LLM审核",
                    "source_issue": "PPT 原文为空，无法对比",
                    "struct_issue": "无",
                }
                sheet.cell(row_index, llm_error_type_column).value = llm_result["error_type"]
                sheet.cell(row_index, llm_severity_column).value = llm_result["severity"]
                sheet.cell(row_index, llm_basis_column).value = llm_result["basis"]
                sheet.cell(row_index, llm_result_column).value = llm_result["result"]
                sheet.cell(row_index, source_issue_column).value = llm_result["source_issue"]
                sheet.cell(row_index, struct_issue_column).value = llm_result["struct_issue"]
                sheet.cell(row_index, root_cause_tag_column).value = "无"
                sheet.cell(row_index, root_cause_reason_column).value = "无"
                processed_rows += 1
                save_if_needed(workbook, excel_path, processed_rows, args.save_batch_size, row_index, str(file_name), page_number)
                continue

            struct_text = build_ppt_struct_text(ppt_struct_obj)
            try:
                llm_result = run_llm_compare(source_text, struct_text)
            except Exception as exc:
                llm_result = {
                    "error_type": "无法判断",
                    "severity": "无法判断",
                    "basis": "LLM 调用失败: {}".format(exc),
                    "result": "未执行LLM审核",
                    "source_issue": "无",
                    "struct_issue": "无",
                }

            sheet.cell(row_index, llm_error_type_column).value = llm_result["error_type"]
            sheet.cell(row_index, llm_severity_column).value = llm_result["severity"]
            sheet.cell(row_index, llm_basis_column).value = llm_result["basis"]
            sheet.cell(row_index, llm_result_column).value = llm_result["result"]
            sheet.cell(row_index, source_issue_column).value = llm_result["source_issue"]
            sheet.cell(row_index, struct_issue_column).value = llm_result["struct_issue"]
        struct_text = build_ppt_struct_text(ppt_struct_obj)
        llm_result_text = str(llm_result["result"] or "")
        if is_empty_llm_result(llm_result_text):
            root_cause_text = "未分析"
            root_cause_reason = "LLM审核结果为空"
        else:
            struct_issue_text = str(llm_result["struct_issue"] or "")
            root_cause = detect_root_cause_by_rules(
                llm_result_text=llm_result_text,
                struct_issue_text=struct_issue_text,
                image_parse_text=image_parse_text,
                chart_parse_text=chart_parse_text,
                paddle_ocr_text=paddle_ocr_text,
                struct_text=struct_text,
            )
            if root_cause is None:
                try:
                    root_cause = call_llm_root_cause(
                        llm_result_text=llm_result_text,
                        struct_issue_text=struct_issue_text,
                        paddle_ocr_text=paddle_ocr_text,
                        image_parse_text=image_parse_text,
                        chart_parse_text=chart_parse_text,
                    )
                except Exception as exc:
                    logger.warning("问题原因分析失败，回退规则归因: row={} file={} page={} error={}", row_index, file_name, page_number, exc)
                    root_cause = build_root_cause_by_terms(
                        struct_issue_text=struct_issue_text,
                        paddle_ocr_text=paddle_ocr_text,
                        image_parse_text=image_parse_text,
                        chart_parse_text=chart_parse_text,
                        struct_text=struct_text,
                    )
            root_cause_text = format_root_cause_text(root_cause["categories"], root_cause["reason"])
            root_cause_reason = format_root_cause_reason(root_cause["categories"], root_cause["reason"])
        sheet.cell(row_index, root_cause_tag_column).value = root_cause_text
        sheet.cell(row_index, root_cause_reason_column).value = root_cause_reason
        processed_rows += 1
        save_if_needed(workbook, excel_path, processed_rows, args.save_batch_size, row_index, str(file_name), page_number)

    workbook.save(excel_path)
    logger.info("Excel 已更新: {}", excel_path)
    logger.info("总处理行数: {}", processed_rows)
    return 0


if __name__ == "__main__":
    sys.exit(main())
