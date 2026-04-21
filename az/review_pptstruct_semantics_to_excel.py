#!/usr/bin/env python3
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

from openpyxl import load_workbook
from PIL import Image, ImageEnhance, ImageOps


STAGE_PATTERN = re.compile(r"(?:[IVX]{1,3}[ABC]?期|[1-4][ABC]?期|STAGE\s*[IVX]{1,3}[ABC]?)", re.IGNORECASE)
TNM_PATTERN = re.compile(r"T\d+[A-Za-z]?\s*N\d+[A-Za-z]?\s*M\d+[A-Za-z]?", re.IGNORECASE)
MUTATION_PATTERNS = [
    re.compile(r"EGFR\s*19[^\n]{0,12}(?:缺失|del|deletion)", re.IGNORECASE),
    re.compile(r"EGFR\s*21[^\n]{0,12}(?:L858R|L858|突变)", re.IGNORECASE),
    re.compile(r"L858R", re.IGNORECASE),
    re.compile(r"T790M", re.IGNORECASE),
    re.compile(r"19del", re.IGNORECASE),
]
DRUG_KEYWORDS = [
    "奥希替尼",
    "吉非替尼",
    "厄洛替尼",
    "埃克替尼",
    "阿美替尼",
    "伏美替尼",
    "奥福替尼",
    "培美曲塞",
    "卡铂",
    "顺铂",
    "帕博利珠单抗",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="基于页截图 OCR 审核 Excel 中的 pptStruct 语义完整性和医学事实错误。")
    parser.add_argument("--excel-path", required=True, help="Excel 文件路径")
    parser.add_argument("--ppt-dir", required=True, help="源 PPT 文件目录")
    parser.add_argument("--sheet-name", help="可选，指定工作表名称；默认使用 active sheet")
    parser.add_argument("--tesseract-lang", default="chi_sim+eng", help="tesseract OCR 语言")
    return parser.parse_args()


def ensure_dependencies() -> None:
    missing = [name for name in ("soffice", "pdftoppm", "pdftotext", "tesseract") if shutil.which(name) is None]
    if missing:
        raise RuntimeError("缺少依赖命令: {}".format(", ".join(missing)))


def normalize_name(name: str) -> str:
    normalized = name.strip().lower()
    normalized = normalized.replace(".pptx.pptx", ".pptx")
    normalized = normalized.replace(".ppt.ppt", ".ppt")
    normalized = re.sub(r"\s+", "", normalized)
    return normalized


def build_ppt_index(ppt_dir: Path) -> Dict[str, Path]:
    index: Dict[str, Path] = {}
    for path in sorted(ppt_dir.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".ppt", ".pptx"}:
            continue
        index[normalize_name(path.name)] = path
    return index


def run_command(command: List[str]) -> None:
    subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)


def render_ppt_to_pdf(ppt_path: Path, work_dir: Path) -> Path:
    pdf_dir = work_dir / "pdf"
    pdf_dir.mkdir(parents=True, exist_ok=True)

    run_command(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_dir),
            str(ppt_path),
        ]
    )
    pdf_path = pdf_dir / "{}.pdf".format(ppt_path.stem)
    if not pdf_path.exists():
        raise RuntimeError("未生成 PDF: {}".format(pdf_path))
    return pdf_path


def render_pdf_page_to_png(pdf_path: Path, page_number: int, work_dir: Path) -> Path:
    img_dir = work_dir / "img"
    img_dir.mkdir(parents=True, exist_ok=True)

    output_prefix = img_dir / "page"
    run_command(
        [
            "pdftoppm",
            "-r",
            "300",
            "-png",
            "-f",
            str(page_number),
            "-l",
            str(page_number),
            "-singlefile",
            str(pdf_path),
            str(output_prefix),
        ]
    )
    png_path = img_dir / "page.png"
    if not png_path.exists():
        raise RuntimeError("未生成页面截图: {}".format(png_path))
    return png_path


def extract_pdf_text(pdf_path: Path, page_number: int) -> str:
    result = subprocess.run(
        [
            "pdftotext",
            "-layout",
            "-f",
            str(page_number),
            "-l",
            str(page_number),
            str(pdf_path),
            "-",
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    return result.stdout


def ocr_image(image_path: Path, lang: str, psm: int = 6) -> str:
    result = subprocess.run(
        ["tesseract", str(image_path), "stdout", "-l", lang, "--psm", str(psm)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    return result.stdout


def preprocess_variants(image_path: Path, work_dir: Path) -> List[Path]:
    base_image = Image.open(image_path).convert("RGB")
    variants: List[Path] = []

    enlarged = base_image.resize((base_image.width * 2, base_image.height * 2))
    variant_specs = [
        ("gray", ImageOps.grayscale(enlarged)),
        ("contrast", ImageEnhance.Contrast(ImageOps.grayscale(enlarged)).enhance(2.2)),
        ("sharp", ImageEnhance.Sharpness(ImageOps.grayscale(enlarged)).enhance(2.0)),
    ]

    for name, image in variant_specs:
        variant_path = work_dir / "{}.png".format(name)
        image.save(variant_path)
        variants.append(variant_path)

    bw_image = ImageEnhance.Contrast(ImageOps.grayscale(enlarged)).enhance(2.8).point(lambda x: 255 if x > 175 else 0)
    bw_path = work_dir / "bw.png"
    bw_image.save(bw_path)
    variants.append(bw_path)
    return variants


def build_tiled_images(image_path: Path, work_dir: Path) -> List[Path]:
    image = Image.open(image_path).convert("RGB")
    width, height = image.size
    tiles: List[Path] = []
    regions = [
        ("top", (0, 0, width, height // 2)),
        ("bottom", (0, height // 2, width, height)),
        ("left", (0, 0, width // 2, height)),
        ("right", (width // 2, 0, width, height)),
    ]
    for name, box in regions:
        tile = image.crop(box).resize((max(1, (box[2] - box[0]) * 2), max(1, (box[3] - box[1]) * 2)))
        tile_path = work_dir / "{}.png".format(name)
        tile.save(tile_path)
        tiles.append(tile_path)
    return tiles


def merge_texts(texts: List[str]) -> str:
    merged_lines: List[str] = []
    seen: Set[str] = set()
    for text in texts:
        for raw_line in text.splitlines():
            line = re.sub(r"\s+", " ", raw_line).strip()
            if len(line) <= 1:
                continue
            key = re.sub(r"\s+", "", line.lower())
            if key in seen:
                continue
            seen.add(key)
            merged_lines.append(line)
    return "\n".join(merged_lines)


def extract_page_text(image_path: Path, pdf_path: Path, page_number: int, lang: str, work_dir: Path) -> str:
    texts: List[str] = [extract_pdf_text(pdf_path, page_number)]
    texts.append(ocr_image(image_path, lang, psm=6))

    variant_paths = preprocess_variants(image_path, work_dir)
    if variant_paths:
        texts.append(ocr_image(variant_paths[1], lang, psm=6))
    tiled_paths = build_tiled_images(image_path, work_dir)
    if tiled_paths:
        texts.append(ocr_image(tiled_paths[0], lang, psm=6))
        texts.append(ocr_image(tiled_paths[1], lang, psm=6))

    return merge_texts(texts)


def collect_texts(value) -> List[str]:
    texts: List[str] = []
    if isinstance(value, dict):
        for child in value.values():
            texts.extend(collect_texts(child))
        text = value.get("text")
        if isinstance(text, str) and text.strip():
            texts.append(text.strip())
        summary = value.get("ppt_summary")
        if isinstance(summary, str) and summary.strip():
            texts.append(summary.strip())
    elif isinstance(value, list):
        for child in value:
            texts.extend(collect_texts(child))
    return texts


def extract_matches(pattern: re.Pattern, text: str) -> Set[str]:
    return {re.sub(r"\s+", "", match.group(0).upper()) for match in pattern.finditer(text)}


def extract_mutations(text: str) -> Set[str]:
    found: Set[str] = set()
    for pattern in MUTATION_PATTERNS:
        for match in pattern.finditer(text):
            found.add(re.sub(r"\s+", "", match.group(0).upper()))
    return found


def extract_drugs(text: str) -> Set[str]:
    return {drug for drug in DRUG_KEYWORDS if drug in text}


def short_join(items: List[str], empty_text: str) -> str:
    if not items:
        return empty_text
    return "；".join(items[:3])


def review_semantics(ocr_text: str, ppt_struct_obj: Dict[str, object]) -> Tuple[str, str, str]:
    ppt_text = "\n".join(collect_texts(ppt_struct_obj))
    ocr_stages = extract_matches(STAGE_PATTERN, ocr_text)
    ppt_stages = extract_matches(STAGE_PATTERN, ppt_text)
    ocr_tnm = extract_matches(TNM_PATTERN, ocr_text)
    ppt_tnm = extract_matches(TNM_PATTERN, ppt_text)
    ocr_mut = extract_mutations(ocr_text)
    ppt_mut = extract_mutations(ppt_text)
    ocr_drugs = extract_drugs(ocr_text)
    ppt_drugs = extract_drugs(ppt_text)

    incomplete: List[str] = []
    factual: List[str] = []

    if ocr_stages and not ppt_stages:
        incomplete.append("未体现原图中的分期：{}".format("、".join(sorted(ocr_stages))))
    elif ocr_stages and ppt_stages and ocr_stages.isdisjoint(ppt_stages):
        factual.append("分期不一致：原图为{}，pptStruct为{}".format("、".join(sorted(ocr_stages)), "、".join(sorted(ppt_stages))))

    if ocr_tnm and not ppt_tnm:
        incomplete.append("未体现原图中的TNM信息：{}".format("、".join(sorted(ocr_tnm))))
    elif ocr_tnm and ppt_tnm and ocr_tnm.isdisjoint(ppt_tnm):
        factual.append("TNM不一致：原图为{}，pptStruct为{}".format("、".join(sorted(ocr_tnm)), "、".join(sorted(ppt_tnm))))

    if ocr_mut and not ppt_mut:
        incomplete.append("未体现原图中的突变信息：{}".format("、".join(sorted(ocr_mut))))
    elif ocr_mut and ppt_mut and ocr_mut.isdisjoint(ppt_mut):
        factual.append("突变信息不一致：原图为{}，pptStruct为{}".format("、".join(sorted(ocr_mut)), "、".join(sorted(ppt_mut))))

    missing_drugs = sorted(ocr_drugs - ppt_drugs)
    extra_drugs = sorted(ppt_drugs - ocr_drugs)
    if missing_drugs:
        incomplete.append("未覆盖原图中的药物信息：{}".format("、".join(missing_drugs)))
    if extra_drugs:
        incomplete.append("额外出现原图未明显识别到的药物信息：{}".format("、".join(extra_drugs)))

    title_text = ""
    ppt_title = ppt_struct_obj.get("ppt_title") if isinstance(ppt_struct_obj, dict) else None
    if isinstance(ppt_title, dict):
        title_text = str(ppt_title.get("text") or "")
    if len(title_text.strip()) <= 4 and (ocr_stages or ocr_mut):
        incomplete.append("标题过于笼统，未体现原图中的关键事实")

    incomplete_text = short_join(incomplete, "未发现明显语义不完整描述")
    factual_text = short_join(factual, "未发现明显医学事实错误")
    if factual:
        conclusion = "存在医学事实风险，优先人工复核"
    elif incomplete:
        conclusion = "主要是语义覆盖不完整"
    else:
        conclusion = "未发现明显问题"
    return incomplete_text, factual_text, conclusion


def main() -> int:
    args = parse_args()
    ensure_dependencies()
    excel_path = Path(args.excel_path).expanduser().resolve()
    ppt_dir = Path(args.ppt_dir).expanduser().resolve()
    workbook = load_workbook(excel_path)
    sheet = workbook[args.sheet_name] if args.sheet_name else workbook.active
    ppt_index = build_ppt_index(ppt_dir)

    sheet.cell(1, 7).value = "语义不完整描述"
    sheet.cell(1, 8).value = "医学事实错误"
    sheet.cell(1, 9).value = "审核结论"
    sheet.column_dimensions["G"].width = 60
    sheet.column_dimensions["H"].width = 60
    sheet.column_dimensions["I"].width = 30
    for column in ("J", "K"):
        sheet.cell(1, ord(column) - 64).value = None

    with tempfile.TemporaryDirectory(prefix="ppt-semantic-review-") as temp_dir:
        temp_root = Path(temp_dir)
        pdf_cache: Dict[Path, Path] = {}
        page_text_cache: Dict[Tuple[Path, int], str] = {}

        for row_index in range(2, sheet.max_row + 1):
            file_name = sheet.cell(row_index, 2).value
            page_number = sheet.cell(row_index, 3).value
            ppt_struct_raw = sheet.cell(row_index, 6).value
            if not file_name or not isinstance(page_number, int) or not ppt_struct_raw:
                continue
            print("processing row={} file={} page={}".format(row_index, file_name, page_number), flush=True)

            ppt_path = ppt_index.get(normalize_name(str(file_name)))
            if not ppt_path:
                sheet.cell(row_index, 7).value = "未找到源PPT文件"
                sheet.cell(row_index, 8).value = "未找到源PPT文件"
                sheet.cell(row_index, 9).value = "无法审核"
                continue

            try:
                ppt_struct_obj = json.loads(str(ppt_struct_raw))
            except Exception as exc:
                sheet.cell(row_index, 7).value = "pptStruct 解析失败: {}".format(exc)
                sheet.cell(row_index, 8).value = "pptStruct 解析失败: {}".format(exc)
                sheet.cell(row_index, 9).value = "无法审核"
                continue

            try:
                pdf_path = pdf_cache.get(ppt_path)
                if pdf_path is None:
                    file_temp_dir = temp_root / re.sub(r"[^A-Za-z0-9._-]+", "_", ppt_path.stem)
                    pdf_path = render_ppt_to_pdf(ppt_path, file_temp_dir)
                    pdf_cache[ppt_path] = pdf_path

                cache_key = (ppt_path, page_number)
                page_text = page_text_cache.get(cache_key)
                if page_text is None:
                    page_temp_dir = temp_root / re.sub(r"[^A-Za-z0-9._-]+", "_", "{}_{}".format(ppt_path.stem, page_number))
                    image_path = render_pdf_page_to_png(pdf_path, page_number, page_temp_dir)
                    page_text = extract_page_text(image_path, pdf_path, page_number, args.tesseract_lang, page_temp_dir)
                    page_text_cache[cache_key] = page_text

                incomplete_text, factual_text, conclusion = review_semantics(page_text, ppt_struct_obj)
            except Exception as exc:
                incomplete_text = "审核失败: {}".format(exc)
                factual_text = "审核失败: {}".format(exc)
                conclusion = "无法审核"

            sheet.cell(row_index, 7).value = incomplete_text
            sheet.cell(row_index, 8).value = factual_text
            sheet.cell(row_index, 9).value = conclusion

    workbook.save(excel_path)
    print("Excel 已更新: {}".format(excel_path))
    return 0


if __name__ == "__main__":
    sys.exit(main())
