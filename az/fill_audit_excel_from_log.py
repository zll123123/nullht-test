#!/usr/bin/env python3
import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="根据 Excel 中已有的报错页码，从日志正文提取对应页的 pptStruct 并回填。")
    parser.add_argument("--excel-path", required=True, help="待回填的 Excel 文件路径")
    parser.add_argument("--log-path", required=True, help="包含 PPT_STRUCT 日志正文的 app 日志路径")
    parser.add_argument("--sheet-name", help="可选，指定工作表名称；默认使用 active sheet")
    parser.add_argument(
        "--skip-non-empty",
        action="store_true",
        help="若第 6 列已有内容则跳过，避免覆盖已写入的 pptStruct",
    )
    return parser.parse_args()


def build_ppt_struct_index(log_path: Path) -> Dict[str, Dict[str, object]]:
    index: Dict[str, Dict[str, object]] = {}
    with log_path.open("r", encoding="utf-8", errors="ignore") as handle:
        for line in handle:
            stripped = line.lstrip()
            if not stripped.startswith('{"1":'):
                continue
            try:
                parsed = json.loads(stripped)
            except json.JSONDecodeError:
                continue
            first_page = parsed.get("1")
            if not isinstance(first_page, dict):
                continue
            first_id = first_page.get("_id")
            if not isinstance(first_id, str) or "_" not in first_id:
                continue
            file_id, _, _ = first_id.rpartition("_")
            if file_id:
                index[file_id] = parsed
    return index


def find_file_id_by_page_map(
    ppt_struct_index: Dict[str, Dict[str, object]], file_name: str, max_page: int
) -> Optional[str]:
    file_name_lower = file_name.lower()
    candidates: List[str] = []
    if "张媛媛" in file_name or "骨转移" in file_name_lower:
        candidates.append("20260415/390f0edabe8d43ff891980329eb1e6b5.pptx")
    if "王杰" in file_name:
        candidates.append("20260415/6d8cc4325c614f939661110d6c9d48cb.pptx")
    if "egfr+nsclc靶向辅助时长及全程管理" in file_name_lower:
        candidates.append("20260331/36a50947636c4c7cb6d450b81478124c.pptx")
    if "刘迎军" in file_name or "ib期egfr阳性nsclc术后辅助治疗病例" in file_name_lower:
        candidates.append("20260415/70c33b5fde334955bdb0091ba1d7c245.pptx")

    for candidate in candidates:
        if candidate in ppt_struct_index:
            return candidate

    for file_id, pages in ppt_struct_index.items():
        if isinstance(pages, dict) and str(max_page) in pages:
            return file_id
    return None


def main() -> int:
    args = parse_args()
    excel_path = Path(args.excel_path).expanduser().resolve()
    log_path = Path(args.log_path).expanduser().resolve()
    workbook = load_workbook(excel_path)
    sheet = workbook[args.sheet_name] if args.sheet_name else workbook.active
    ppt_struct_index = build_ppt_struct_index(log_path)

    file_id_cache: Dict[str, Optional[str]] = {}
    updated_rows = 0

    for row_index in range(2, sheet.max_row + 1):
        file_name = sheet.cell(row_index, 2).value
        page_number = sheet.cell(row_index, 3).value
        existing_value = sheet.cell(row_index, 6).value
        if args.skip_non_empty and existing_value:
            continue
        if not file_name or not isinstance(page_number, int):
            continue

        file_name = str(file_name)
        if file_name not in file_id_cache:
            file_id_cache[file_name] = find_file_id_by_page_map(ppt_struct_index, file_name, page_number)
        file_id = file_id_cache[file_name]
        if not file_id:
            sheet.cell(row_index, 6).value = "日志中未定位到文件对应的PPT_STRUCT"
            updated_rows += 1
            continue

        page_map = ppt_struct_index.get(file_id) or {}
        page_object = page_map.get(str(page_number)) if isinstance(page_map, dict) else None
        sheet.cell(row_index, 6).value = (
            json.dumps(page_object, ensure_ascii=False) if page_object is not None else "日志中未定位到该页pptStruct"
        )
        updated_rows += 1

    workbook.save(excel_path)
    print(f"Excel 已更新: {excel_path}")
    print(f"写入记录数: {updated_rows}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
